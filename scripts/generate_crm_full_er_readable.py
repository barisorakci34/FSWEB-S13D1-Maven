#!/usr/bin/env python3
"""Generate readable ER package for full CRM workbook."""

from __future__ import annotations

import argparse
import csv
import math
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Sequence, Tuple

import matplotlib.pyplot as plt
import networkx as nx
from openpyxl import load_workbook


@dataclass
class FkRow:
    constraint: str
    source_table: str
    source_columns: List[str]
    target_table: str


AUDIT_USER_COLUMN_RE = re.compile(r"^(CreatedBy|UpdatedBy|DeletedBy)(UserUid)?$", re.IGNORECASE)


def short_name(full_name: str) -> str:
    return full_name.split(".", 1)[1] if "." in full_name else full_name


def slug(text: str) -> str:
    s = text.lower().replace(".", "-")
    s = re.sub(r"[^a-z0-9\\-]+", "-", s)
    return re.sub(r"-+", "-", s).strip("-")


def split_cols(value: str) -> List[str]:
    if not value:
        return []
    return [x.strip() for x in str(value).split(",") if x.strip()]


def summarize_columns(columns: Iterable[str], max_show: int = 2) -> str:
    uniq: List[str] = []
    seen = set()
    for col in columns:
        key = col.lower()
        if key in seen:
            continue
        seen.add(key)
        uniq.append(col)
    if not uniq:
        return ""
    if len(uniq) <= max_show:
        return ", ".join(uniq)
    return ", ".join(uniq[:max_show]) + f" +{len(uniq)-max_show}"


def read_workbook(input_xlsx: Path) -> Tuple[Dict[str, int], List[FkRow]]:
    wb = load_workbook(input_xlsx, read_only=True, data_only=True)
    ws_tables = wb["Tablolar"]
    ws_fk = wb["Iliskiler_FK"]

    table_col_counts: Dict[str, int] = {}
    for row in ws_tables.iter_rows(min_row=2, values_only=True):
        schema, table, col_count = row[0], row[1], row[2]
        if not (schema and table):
            continue
        table_col_counts[f"{schema}.{table}"] = int(col_count or 0)

    fk_rows: List[FkRow] = []
    for row in ws_fk.iter_rows(min_row=2, values_only=True):
        constraint, ss, st, sc, ts, tt = row[0], row[1], row[2], row[3], row[4], row[5]
        if not (ss and st and ts and tt):
            continue
        fk_rows.append(
            FkRow(
                constraint=str(constraint or ""),
                source_table=f"{ss}.{st}",
                source_columns=split_cols(str(sc or "")),
                target_table=f"{ts}.{tt}",
            )
        )
    return table_col_counts, fk_rows


def is_audit_user_fk(row: FkRow) -> bool:
    if row.target_table != "Crm.Users":
        return False
    if not row.source_columns:
        return False
    return all(AUDIT_USER_COLUMN_RE.match(c) for c in row.source_columns)


def aggregate_fk_pairs(rows: Sequence[FkRow]) -> Dict[Tuple[str, str], Dict[str, object]]:
    pairs: Dict[Tuple[str, str], Dict[str, object]] = {}
    for row in rows:
        key = (row.source_table, row.target_table)
        if key not in pairs:
            pairs[key] = {"count": 0, "columns": [], "constraints": []}
        pairs[key]["count"] = int(pairs[key]["count"]) + 1
        pairs[key]["columns"].extend(row.source_columns)
        pairs[key]["constraints"].append(row.constraint)
    return dict(sorted(pairs.items(), key=lambda kv: (kv[0][0], kv[0][1])))


def build_graph(
    pair_map: Dict[Tuple[str, str], Dict[str, object]],
    table_col_counts: Dict[str, int],
) -> nx.DiGraph:
    g = nx.DiGraph()
    nodes = set(table_col_counts.keys())
    for src, dst in pair_map.keys():
        nodes.add(src)
        nodes.add(dst)
    for n in nodes:
        g.add_node(n, col_count=table_col_counts.get(n, 0))

    for (src, dst), payload in pair_map.items():
        g.add_edge(
            src,
            dst,
            fk_count=int(payload["count"]),
            col_label=summarize_columns(payload["columns"], max_show=2),
        )
    return g


def compute_layout(g: nx.DiGraph, seed: int) -> Dict[str, Tuple[float, float]]:
    if g.number_of_nodes() == 0:
        return {}
    k = 1.65 / math.sqrt(max(1, g.number_of_nodes()))
    return nx.spring_layout(g.to_undirected(), seed=seed, k=k, iterations=800)


def draw_graph(
    g: nx.DiGraph,
    output_base: Path,
    title: str,
    label_limit: int | None = None,
    show_edge_labels: bool = False,
    seed: int = 42,
) -> None:
    pos = compute_layout(g, seed=seed)
    fig, ax = plt.subplots(figsize=(30, 20))
    ax.set_axis_off()

    indeg = dict(g.in_degree())
    outdeg = dict(g.out_degree())
    degree = {n: indeg.get(n, 0) + outdeg.get(n, 0) for n in g.nodes}

    node_sizes = []
    node_colors = []
    for n in g.nodes:
        d = degree.get(n, 0)
        c = g.nodes[n].get("col_count", 0)
        node_sizes.append(130 + min(2400, d * 38 + c * 0.55))
        if n == "Crm.Users":
            node_colors.append("#D9534F")
        elif d >= 30:
            node_colors.append("#F0AD4E")
        elif d >= 10:
            node_colors.append("#5BC0DE")
        else:
            node_colors.append("#9FBAD0")

    edge_widths = [0.45 + min(3.0, g.edges[e]["fk_count"] * 0.26) for e in g.edges]
    nx.draw_networkx_edges(
        g,
        pos=pos,
        arrows=True,
        arrowstyle="-|>",
        arrowsize=10,
        width=edge_widths,
        alpha=0.22,
        edge_color="#2F3A45",
        connectionstyle="arc3,rad=0.03",
        ax=ax,
    )
    nx.draw_networkx_nodes(
        g,
        pos=pos,
        node_size=node_sizes,
        node_color=node_colors,
        edgecolors="#1F2A35",
        linewidths=0.5,
        ax=ax,
    )

    label_nodes: Sequence[str]
    if label_limit is None or label_limit >= g.number_of_nodes():
        label_nodes = list(g.nodes)
    else:
        label_nodes = [n for n, _ in sorted(degree.items(), key=lambda kv: (-kv[1], kv[0]))[:label_limit]]
    labels = {n: short_name(n) for n in label_nodes}
    nx.draw_networkx_labels(g, pos=pos, labels=labels, font_size=6.2, font_color="#101418", ax=ax)

    if show_edge_labels:
        edge_labels = {
            (u, v): d.get("col_label", "")
            for u, v, d in g.edges(data=True)
            if d.get("col_label", "")
        }
        if edge_labels:
            nx.draw_networkx_edge_labels(
                g,
                pos=pos,
                edge_labels=edge_labels,
                font_size=5.6,
                rotate=False,
                bbox={"facecolor": "white", "edgecolor": "none", "alpha": 0.78, "pad": 0.12},
                ax=ax,
            )

    ax.set_title(
        (
            f"{title}\n"
            f"Node: {g.number_of_nodes()} | Aggregated edges: {g.number_of_edges()} | Labelled: {len(labels)}"
        ),
        fontsize=14,
        pad=14,
    )

    legend = [
        plt.Line2D([0], [0], marker="o", color="w", label="High degree (>=30)", markerfacecolor="#F0AD4E", markeredgecolor="#1F2A35", markersize=8),
        plt.Line2D([0], [0], marker="o", color="w", label="Medium degree (>=10)", markerfacecolor="#5BC0DE", markeredgecolor="#1F2A35", markersize=8),
        plt.Line2D([0], [0], marker="o", color="w", label="Other nodes", markerfacecolor="#9FBAD0", markeredgecolor="#1F2A35", markersize=8),
        plt.Line2D([0], [0], color="#2F3A45", lw=2, label="Explicit FK"),
    ]
    ax.legend(handles=legend, loc="upper right", framealpha=0.92, fontsize=9)

    output_base.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output_base.with_suffix(".png"), dpi=420, format="png", bbox_inches="tight", pad_inches=0.2)
    fig.savefig(output_base.with_suffix(".svg"), format="svg", bbox_inches="tight", pad_inches=0.2)
    plt.close(fig)


def draw_audit_users_pages(
    audit_rows: Sequence[FkRow],
    table_col_counts: Dict[str, int],
    output_dir: Path,
    page_size: int = 90,
) -> int:
    source_map: Dict[str, List[FkRow]] = defaultdict(list)
    for row in audit_rows:
        source_map[row.source_table].append(row)
    sources = sorted(source_map.keys(), key=lambda s: (-len(source_map[s]), s))
    if not sources:
        return 0

    pages = math.ceil(len(sources) / page_size)
    for page in range(pages):
        page_sources = sources[page * page_size : (page + 1) * page_size]
        graph = nx.DiGraph()
        graph.add_node("Crm.Users")
        for src in page_sources:
            graph.add_node(src)
            graph.add_edge(src, "Crm.Users", fk_count=len(source_map[src]))

        radius = max(8.0, len(page_sources) * 0.17)
        pos = {"Crm.Users": (0.0, 0.0)}
        for idx, src in enumerate(page_sources):
            angle = (2.0 * math.pi * idx) / max(1, len(page_sources))
            pos[src] = (radius * math.cos(angle), radius * math.sin(angle))

        fig, ax = plt.subplots(figsize=(24, 24))
        ax.set_axis_off()

        nodes = ["Crm.Users"] + page_sources
        sizes = [4300] + [650 + min(1800, table_col_counts.get(src, 0) * 0.8 + len(source_map[src]) * 120) for src in page_sources]
        colors = ["#D9534F"] + ["#A9CCE3" for _ in page_sources]

        nx.draw_networkx_nodes(
            graph,
            pos=pos,
            nodelist=nodes,
            node_size=sizes,
            node_color=colors,
            edgecolors="#1F2A35",
            linewidths=1.0,
            ax=ax,
        )
        nx.draw_networkx_edges(
            graph,
            pos=pos,
            arrows=True,
            arrowstyle="-|>",
            arrowsize=16,
            width=1.25,
            alpha=0.68,
            edge_color="#2F3A45",
            connectionstyle="arc3,rad=0.06",
            ax=ax,
        )

        labels = {n: short_name(n) for n in nodes}
        nx.draw_networkx_labels(graph, pos=pos, labels=labels, font_size=7.8, font_color="#101418", ax=ax)
        nx.draw_networkx_labels(
            graph,
            pos={"Crm.Users": pos["Crm.Users"]},
            labels={"Crm.Users": "Users"},
            font_size=11.0,
            font_color="#FFFFFF",
            font_weight="bold",
            ax=ax,
        )

        ax.set_title(
            (
                f"CRM Audit FK to Crm.Users (page {page+1}/{pages})\n"
                f"Page source tables: {len(page_sources)} | Total source tables: {len(sources)}"
            ),
            fontsize=14,
            pad=14,
        )

        out = output_dir / f"audit-users-p{page+1}-highres"
        fig.savefig(out.with_suffix(".png"), dpi=420, format="png", bbox_inches="tight", pad_inches=0.2)
        fig.savefig(out.with_suffix(".svg"), format="svg", bbox_inches="tight", pad_inches=0.2)
        plt.close(fig)

    return pages


def draw_community_diagrams(
    business_graph: nx.DiGraph,
    output_dir: Path,
) -> List[Tuple[int, int, int, str]]:
    undirected = business_graph.to_undirected()
    communities = list(nx.algorithms.community.greedy_modularity_communities(undirected, weight="fk_count"))
    communities = sorted(communities, key=lambda s: (-len(s), sorted(list(s))[0] if s else ""))

    summary_rows: List[Tuple[int, int, int, str]] = []
    for idx, community_nodes in enumerate(communities, start=1):
        nodes = sorted(list(community_nodes))
        if len(nodes) < 2:
            continue
        sub = business_graph.subgraph(nodes).copy()
        if sub.number_of_edges() == 0:
            continue

        label_limit = None if len(nodes) <= 35 else 35
        show_edge_labels = len(nodes) <= 26 and sub.number_of_edges() <= 80
        seed = 100 + idx
        title = f"CRM Business ER Community {idx}"
        out = output_dir / f"community-{idx:02d}-highres"
        draw_graph(sub, out, title=title, label_limit=label_limit, show_edge_labels=show_edge_labels, seed=seed)

        top_node = max(nodes, key=lambda n: sub.degree(n))
        summary_rows.append((idx, len(nodes), sub.number_of_edges(), top_node))

    return summary_rows


def write_csv(
    output_path: Path,
    headers: Sequence[str],
    rows: Sequence[Sequence[object]],
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding="utf-8") as f:
        wr = csv.writer(f)
        wr.writerow(headers)
        wr.writerows(rows)


def write_readme(
    output_dir: Path,
    total_fk: int,
    audit_fk: int,
    business_fk: int,
    business_nodes: int,
    business_edges: int,
    community_count: int,
    audit_pages: int,
) -> None:
    readme = output_dir / "README.md"
    lines = [
        "# CRM Full ER - Readable Package",
        "",
        "Bu paket, onceki karisik tek-gorsel yerine katmanli ve okunabilir cikti verir.",
        "",
        "## Ozet",
        "",
        f"- Toplam FK: **{total_fk}**",
        f"- Audit FK (Crm.Users + CreatedBy/UpdatedBy/DeletedBy): **{audit_fk}**",
        f"- Business FK (kalan): **{business_fk}**",
        f"- Business graph node: **{business_nodes}**",
        f"- Business graph aggregated edge: **{business_edges}**",
        f"- Community diyagram sayisi: **{community_count}**",
        f"- Audit Users sayfa sayisi: **{audit_pages}**",
        "",
        "## Dosyalar",
        "",
        "- `00-business-overview-highres.png/.svg`",
        "- `01-business-overview-without-users-highres.png/.svg`",
        "- `community-XX-highres.png/.svg` (modul bazli kume diyagramlari)",
        "- `audit-users-pX-highres.png/.svg` (audit FK'ler ayri)",
        "- `fk-business-pairs.csv`",
        "- `fk-audit-users.csv`",
        "- `communities-summary.csv`",
        "",
        "Okuma sirasi: once 00 -> 01 -> community dosyalari -> audit-users sayfalari.",
        "",
    ]
    readme.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate readable full CRM ER package")
    parser.add_argument(
        "--input-xlsx",
        default="docs/crm-full-veri-sozlugu.xlsx",
        help="Input workbook path",
    )
    parser.add_argument(
        "--output-dir",
        default="docs/crm-full-er-readable",
        help="Output directory",
    )
    parser.add_argument(
        "--audit-page-size",
        type=int,
        default=90,
        help="Number of source tables per audit-users page",
    )
    args = parser.parse_args()

    input_xlsx = Path(args.input_xlsx)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    table_col_counts, fk_rows = read_workbook(input_xlsx)
    audit_rows = [r for r in fk_rows if is_audit_user_fk(r)]
    business_rows = [r for r in fk_rows if not is_audit_user_fk(r)]

    business_pairs = aggregate_fk_pairs(business_rows)
    business_graph = build_graph(business_pairs, table_col_counts)

    draw_graph(
        business_graph,
        output_base=output_dir / "00-business-overview-highres",
        title="CRM Business ER (Audit FK removed)",
        label_limit=85,
        show_edge_labels=False,
        seed=41,
    )

    business_without_users = business_graph.copy()
    if "Crm.Users" in business_without_users:
        business_without_users.remove_node("Crm.Users")
    draw_graph(
        business_without_users,
        output_base=output_dir / "01-business-overview-without-users-highres",
        title="CRM Business ER (without Crm.Users)",
        label_limit=95,
        show_edge_labels=False,
        seed=43,
    )

    community_summary = draw_community_diagrams(business_graph, output_dir=output_dir)
    audit_pages = draw_audit_users_pages(
        audit_rows=audit_rows,
        table_col_counts=table_col_counts,
        output_dir=output_dir,
        page_size=max(20, args.audit_page_size),
    )

    business_rows_csv = []
    for (src, dst), payload in business_pairs.items():
        business_rows_csv.append([src, dst, int(payload["count"]), summarize_columns(payload["columns"], max_show=4)])
    write_csv(
        output_path=output_dir / "fk-business-pairs.csv",
        headers=["source_table", "target_table", "fk_count", "sample_source_columns"],
        rows=business_rows_csv,
    )

    audit_rows_csv = []
    for r in audit_rows:
        audit_rows_csv.append([r.constraint, r.source_table, ",".join(r.source_columns), r.target_table])
    write_csv(
        output_path=output_dir / "fk-audit-users.csv",
        headers=["constraint", "source_table", "source_columns", "target_table"],
        rows=audit_rows_csv,
    )

    write_csv(
        output_path=output_dir / "communities-summary.csv",
        headers=["community_id", "node_count", "edge_count", "top_degree_node"],
        rows=community_summary,
    )

    write_readme(
        output_dir=output_dir,
        total_fk=len(fk_rows),
        audit_fk=len(audit_rows),
        business_fk=len(business_rows),
        business_nodes=business_graph.number_of_nodes(),
        business_edges=business_graph.number_of_edges(),
        community_count=len(community_summary),
        audit_pages=audit_pages,
    )

    print(f"Input: {input_xlsx}")
    print(f"Output dir: {output_dir}")
    print(f"Total FK: {len(fk_rows)}")
    print(f"Audit FK: {len(audit_rows)}")
    print(f"Business FK: {len(business_rows)}")
    print(f"Business nodes: {business_graph.number_of_nodes()}")
    print(f"Business aggregated edges: {business_graph.number_of_edges()}")
    print(f"Community diagrams: {len(community_summary)}")
    print(f"Audit users pages: {audit_pages}")


if __name__ == "__main__":
    main()
