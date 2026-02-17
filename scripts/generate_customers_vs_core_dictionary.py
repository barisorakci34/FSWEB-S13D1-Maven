#!/usr/bin/env python3
"""Generate data dictionary and relationship workbook from SQL script."""

from __future__ import annotations

import argparse
import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


@dataclass
class ColumnInfo:
    name: str
    data_type: str
    nullable: str
    is_computed: bool = False
    computed_expression: str = ""
    is_masked: bool = False
    is_identity: bool = False


@dataclass
class UniqueConstraint:
    name: str
    columns: List[str] = field(default_factory=list)


@dataclass
class ForeignKeyInfo:
    constraint: str
    source_schema: str
    source_table: str
    source_columns: List[str]
    target_schema: str
    target_table: str
    target_columns: List[str]


@dataclass
class TableInfo:
    schema: str
    table: str
    columns: List[ColumnInfo] = field(default_factory=list)
    pk_constraint: str = ""
    pk_columns: List[str] = field(default_factory=list)
    unique_constraints: List[UniqueConstraint] = field(default_factory=list)

    @property
    def full_name(self) -> str:
        return f"{self.schema}.{self.table}"


def split_top_level(source: str) -> List[str]:
    parts: List[str] = []
    current: List[str] = []
    depth = 0
    for ch in source:
        if ch == "(":
            depth += 1
        elif ch == ")" and depth > 0:
            depth -= 1
        if ch == "," and depth == 0:
            item = "".join(current).strip()
            if item:
                parts.append(item)
            current = []
        else:
            current.append(ch)
    tail = "".join(current).strip()
    if tail:
        parts.append(tail)
    return parts


def extract_bracket_columns(segment: str) -> List[str]:
    return [x.strip() for x in re.findall(r"\[([^\]]+)\]", segment, flags=re.IGNORECASE)]


def parse_column(entry: str) -> ColumnInfo | None:
    line = " ".join(entry.replace("\r", " ").replace("\n", " ").split())

    computed = re.match(r"^\[(?P<name>[^\]]+)\]\s+AS\s+(?P<expr>.+)$", line, flags=re.IGNORECASE)
    if computed:
        return ColumnInfo(
            name=computed.group("name"),
            data_type="COMPUTED",
            nullable="N/A",
            is_computed=True,
            computed_expression=computed.group("expr").strip(),
            is_masked=False,
            is_identity=False,
        )

    m = re.match(r"^\[(?P<name>[^\]]+)\]\s+\[(?P<dtype>[^\]]+)\](?P<rest>.*)$", line, flags=re.IGNORECASE)
    if not m:
        return None

    name = m.group("name").strip()
    dtype = m.group("dtype").strip()
    rest = m.group("rest").strip()

    type_size = ""
    if rest.startswith("("):
        d = 0
        for idx, ch in enumerate(rest):
            if ch == "(":
                d += 1
            elif ch == ")":
                d -= 1
                if d == 0:
                    type_size = rest[: idx + 1]
                    rest = rest[idx + 1 :].strip()
                    break

    data_type = f"{dtype}{type_size}"
    nullable = "UNKNOWN"
    if re.search(r"\bNOT\s+NULL\b", rest, flags=re.IGNORECASE):
        nullable = "NOT NULL"
    elif re.search(r"\bNULL\b", rest, flags=re.IGNORECASE):
        nullable = "NULL"

    is_masked = bool(re.search(r"\bMASKED\s+WITH\b", rest, flags=re.IGNORECASE))
    is_identity = bool(re.search(r"\bIDENTITY\s*\(", rest, flags=re.IGNORECASE))

    return ColumnInfo(
        name=name,
        data_type=data_type,
        nullable=nullable,
        is_computed=False,
        computed_expression="",
        is_masked=is_masked,
        is_identity=is_identity,
    )


def parse_tables(sql_text: str) -> Dict[str, TableInfo]:
    pattern = re.compile(
        r"CREATE\s+TABLE\s+\[(?P<schema>[^\]]+)\]\.\[(?P<table>[^\]]+)\]\((?P<body>.*?)\)\s*ON\s+\[PRIMARY\](?:\s*TEXTIMAGE_ON\s+\[PRIMARY\])?",
        flags=re.IGNORECASE | re.DOTALL,
    )

    tables: Dict[str, TableInfo] = {}
    for match in pattern.finditer(sql_text):
        schema = match.group("schema").strip()
        table = match.group("table").strip()
        body = match.group("body")
        key = f"{schema}.{table}"

        info = TableInfo(schema=schema, table=table)
        entries = split_top_level(body)

        for entry in entries:
            entry_strip = entry.strip()
            if entry_strip.upper().startswith("CONSTRAINT"):
                pk = re.search(
                    r"CONSTRAINT\s+\[(?P<name>[^\]]+)\]\s+PRIMARY\s+KEY\b[^(]*\((?P<cols>.*?)\)",
                    entry_strip,
                    flags=re.IGNORECASE | re.DOTALL,
                )
                if pk:
                    info.pk_constraint = pk.group("name").strip()
                    info.pk_columns = extract_bracket_columns(pk.group("cols"))
                    continue

                uq = re.search(
                    r"CONSTRAINT\s+\[(?P<name>[^\]]+)\]\s+UNIQUE\b[^(]*\((?P<cols>.*?)\)",
                    entry_strip,
                    flags=re.IGNORECASE | re.DOTALL,
                )
                if uq:
                    info.unique_constraints.append(
                        UniqueConstraint(
                            name=uq.group("name").strip(),
                            columns=extract_bracket_columns(uq.group("cols")),
                        )
                    )
                continue

            col = parse_column(entry_strip)
            if col:
                info.columns.append(col)

        tables[key] = info

    return dict(sorted(tables.items(), key=lambda kv: kv[0].lower()))


def parse_defaults(sql_text: str) -> Dict[Tuple[str, str], Tuple[str, str]]:
    pattern = re.compile(
        r"ALTER\s+TABLE\s+\[(?P<schema>[^\]]+)\]\.\[(?P<table>[^\]]+)\]\s+ADD\s+(?:CONSTRAINT\s+\[(?P<constraint>[^\]]+)\]\s+)?DEFAULT\s+(?P<expr>.+?)\s+FOR\s+\[(?P<column>[^\]]+)\]",
        flags=re.IGNORECASE,
    )
    defaults: Dict[Tuple[str, str], Tuple[str, str]] = {}
    for m in pattern.finditer(sql_text):
        full_table = f"{m.group('schema').strip()}.{m.group('table').strip()}"
        column = m.group("column").strip()
        expr = m.group("expr").strip()
        constraint_name = (m.group("constraint") or "").strip()
        defaults[(full_table, column)] = (expr, constraint_name)
    return defaults


def parse_foreign_keys(sql_text: str) -> List[ForeignKeyInfo]:
    pattern = re.compile(
        r"ALTER\s+TABLE\s+\[(?P<src_schema>[^\]]+)\]\.\[(?P<src_table>[^\]]+)\]\s+WITH\s+CHECK\s+ADD\s+CONSTRAINT\s+\[(?P<constraint>[^\]]+)\]\s+FOREIGN\s+KEY\((?P<src_cols>.*?)\)\s*REFERENCES\s+\[(?P<ref_schema>[^\]]+)\]\.\[(?P<ref_table>[^\]]+)\]\s*\((?P<ref_cols>.*?)\)",
        flags=re.IGNORECASE | re.DOTALL,
    )
    fks: List[ForeignKeyInfo] = []
    for m in pattern.finditer(sql_text):
        fks.append(
            ForeignKeyInfo(
                constraint=m.group("constraint").strip(),
                source_schema=m.group("src_schema").strip(),
                source_table=m.group("src_table").strip(),
                source_columns=extract_bracket_columns(m.group("src_cols")),
                target_schema=m.group("ref_schema").strip(),
                target_table=m.group("ref_table").strip(),
                target_columns=extract_bracket_columns(m.group("ref_cols")),
            )
        )
    return sorted(fks, key=lambda x: (x.source_schema, x.source_table, x.constraint))


def infer_relationships(
    tables: Dict[str, TableInfo], explicit_fks: Sequence[ForeignKeyInfo]
) -> List[Tuple[str, str, str, str, str]]:
    explicit_pairs = set()
    for fk in explicit_fks:
        source = f"{fk.source_schema}.{fk.source_table}"
        for col in fk.source_columns:
            explicit_pairs.add((source, col.upper()))

    table_names = {}
    for full_name, tbl in tables.items():
        t = tbl.table
        lower = t.lower()
        singular = lower[:-1] if lower.endswith("s") else lower
        table_names[full_name] = (lower, singular)

    results = []
    seen = set()
    for source_full, tbl in tables.items():
        for col in tbl.columns:
            if col.is_computed:
                continue
            name = col.name
            if not name.lower().endswith("uid") or name.lower() == "uid":
                continue
            if (source_full, name.upper()) in explicit_pairs:
                continue

            base = name[:-3].lower()
            candidates: List[Tuple[str, str]] = []
            for target_full, (tbl_lower, singular) in table_names.items():
                if target_full == source_full:
                    if base in {singular, tbl_lower}:
                        candidates.append((target_full, "high"))
                    elif base.endswith(singular):
                        candidates.append((target_full, "medium"))
                    continue

                if base in {singular, tbl_lower}:
                    candidates.append((target_full, "high"))
                elif base.endswith(singular):
                    candidates.append((target_full, "medium"))

            if len(candidates) != 1:
                continue

            target, confidence = candidates[0]
            key = (source_full, name, target)
            if key in seen:
                continue
            seen.add(key)
            reason = "Uid kolon ismi tablo ismiyle eslesiyor"
            results.append((source_full, name, target, confidence, reason))

    return sorted(results, key=lambda x: (x[0], x[2], x[1]))


def style_headers(ws, widths: List[int]) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def generate_workbook(
    source_file: Path,
    output_file: Path,
    tables: Dict[str, TableInfo],
    defaults: Dict[Tuple[str, str], Tuple[str, str]],
    foreign_keys: Sequence[ForeignKeyInfo],
    inferred: Sequence[Tuple[str, str, str, str, str]],
) -> None:
    wb = Workbook()
    ws_ozet = wb.active
    ws_ozet.title = "Ozet"
    ws_tables = wb.create_sheet("Tablolar")
    ws_cols = wb.create_sheet("Kolonlar")
    ws_fk = wb.create_sheet("Iliskiler_FK")
    ws_guess = wb.create_sheet("Iliskiler_Tahmini")

    fk_out = defaultdict(int)
    fk_in = defaultdict(int)
    for fk in foreign_keys:
        src = f"{fk.source_schema}.{fk.source_table}"
        tgt = f"{fk.target_schema}.{fk.target_table}"
        fk_out[src] += 1
        fk_in[tgt] += 1

    total_columns = sum(len(t.columns) for t in tables.values())
    total_uniques = sum(len(t.unique_constraints) for t in tables.values())

    ws_ozet.append(["Alan", "Deger"])
    ws_ozet.append(["Kaynak Dosya", str(source_file)])
    ws_ozet.append(["Uretim Zamanı", datetime.now().isoformat(timespec="seconds")])
    ws_ozet.append(["Toplam Tablo", len(tables)])
    ws_ozet.append(["Toplam Kolon", total_columns])
    ws_ozet.append(["Toplam PK", sum(1 for t in tables.values() if t.pk_columns)])
    ws_ozet.append(["Toplam Unique Constraint", total_uniques])
    ws_ozet.append(["Toplam Default Constraint", len(defaults)])
    ws_ozet.append(["Toplam FK", len(foreign_keys)])
    ws_ozet.append(["Toplam Tahmini Iliski", len(inferred)])
    style_headers(ws_ozet, [35, 80])

    ws_tables.append(
        [
            "Schema",
            "Table",
            "ColumnCount",
            "PKConstraint",
            "PKColumns",
            "UniqueConstraintCount",
            "DefaultCount",
            "FKOutCount",
            "FKInCount",
        ]
    )
    for full, t in tables.items():
        default_count = sum(1 for (tbl, _), _ in defaults.items() if tbl == full)
        ws_tables.append(
            [
                t.schema,
                t.table,
                len(t.columns),
                t.pk_constraint,
                ", ".join(t.pk_columns),
                len(t.unique_constraints),
                default_count,
                fk_out.get(full, 0),
                fk_in.get(full, 0),
            ]
        )
    style_headers(ws_tables, [14, 26, 12, 34, 26, 20, 14, 12, 12])

    ws_cols.append(
        [
            "Schema",
            "Table",
            "Ordinal",
            "Column",
            "DataType",
            "Nullability",
            "IsPrimaryKey",
            "IsUnique",
            "IsComputed",
            "ComputedExpression",
            "IsMasked",
            "IsIdentity",
            "DefaultExpression",
            "DefaultConstraint",
        ]
    )
    for full, t in tables.items():
        pk_set = {c.upper() for c in t.pk_columns}
        unique_cols = {
            c.upper()
            for uq in t.unique_constraints
            for c in uq.columns
        }
        for idx, c in enumerate(t.columns, start=1):
            default_expr, default_constraint = defaults.get((full, c.name), ("", ""))
            ws_cols.append(
                [
                    t.schema,
                    t.table,
                    idx,
                    c.name,
                    c.data_type,
                    c.nullable,
                    "Yes" if c.name.upper() in pk_set else "No",
                    "Yes" if c.name.upper() in unique_cols else "No",
                    "Yes" if c.is_computed else "No",
                    c.computed_expression,
                    "Yes" if c.is_masked else "No",
                    "Yes" if c.is_identity else "No",
                    default_expr,
                    default_constraint,
                ]
            )
    style_headers(ws_cols, [12, 24, 10, 30, 20, 12, 12, 10, 11, 55, 10, 10, 30, 36])

    ws_fk.append(
        [
            "Constraint",
            "SourceSchema",
            "SourceTable",
            "SourceColumns",
            "TargetSchema",
            "TargetTable",
            "TargetColumns",
            "TargetInScript",
            "Cardinality",
        ]
    )
    table_keys = set(tables.keys())
    for fk in foreign_keys:
        target_full = f"{fk.target_schema}.{fk.target_table}"
        ws_fk.append(
            [
                fk.constraint,
                fk.source_schema,
                fk.source_table,
                ", ".join(fk.source_columns),
                fk.target_schema,
                fk.target_table,
                ", ".join(fk.target_columns),
                "Yes" if target_full in table_keys else "No",
                "N:1",
            ]
        )
    style_headers(ws_fk, [44, 14, 22, 24, 14, 22, 24, 14, 12])

    ws_guess.append(
        [
            "SourceTable",
            "SourceColumn",
            "TargetTable",
            "Confidence",
            "Reason",
        ]
    )
    for source, col, target, confidence, reason in inferred:
        ws_guess.append([source, col, target, confidence, reason])
    style_headers(ws_guess, [30, 24, 30, 14, 46])

    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_file)


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate customers-vs-core dictionary xlsx")
    parser.add_argument(
        "--input",
        default="/home/ubuntu/.cursor/projects/workspace/uploads/customers-vs-core.txt",
        help="Input SQL script path",
    )
    parser.add_argument(
        "--output",
        default="docs/customers-vs-core-veri-sozlugu.xlsx",
        help="Output XLSX path",
    )
    args = parser.parse_args()

    source_file = Path(args.input)
    output_file = Path(args.output)
    if not source_file.exists():
        raise FileNotFoundError(f"Input file not found: {source_file}")

    sql_text = source_file.read_text(encoding="utf-8")
    tables = parse_tables(sql_text)
    defaults = parse_defaults(sql_text)
    foreign_keys = parse_foreign_keys(sql_text)
    inferred = infer_relationships(tables, foreign_keys)

    generate_workbook(
        source_file=source_file,
        output_file=output_file,
        tables=tables,
        defaults=defaults,
        foreign_keys=foreign_keys,
        inferred=inferred,
    )

    print(f"Input: {source_file}")
    print(f"Output: {output_file}")
    print(f"Tables: {len(tables)}")
    print(f"Columns: {sum(len(t.columns) for t in tables.values())}")
    print(f"Defaults: {len(defaults)}")
    print(f"FKs: {len(foreign_keys)}")
    print(f"Inferred relationships: {len(inferred)}")


if __name__ == "__main__":
    main()
