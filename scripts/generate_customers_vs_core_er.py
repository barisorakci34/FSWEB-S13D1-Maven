#!/usr/bin/env python3
"""Generate ER diagrams from customers-vs-core dictionary workbook."""

from __future__ import annotations

import argparse
from collections import defaultdict
from pathlib import Path
from typing import Dict, Iterable, List, Sequence, Tuple

import matplotlib.pyplot as plt
import networkx as nx
from openpyxl import load_workbook

FkEdge = Tuple[str, str, str, bool, str]  # src, dst, src_cols, target_in_script, constraint
InferredEdge = Tuple[str, str, str, str]  # src, dst, src_col, confidence


def short_name(full_name: str) -> str:
    return full_name.split(".", 1)[1] if "." in full_name else full_name


def compress_columns(columns: Iterable[str], max_show: int = 2) -> str:
    vals = [c.strip() for c in columns if c and c.strip()]
    if not vals:
        return ""
    unique_vals: List[str] = []
    seen = set()
    for col in vals:
        key = col.lower()
        if key in seen:
            continue
        seen.add(key)
        unique_vals.append(col)
    if len(unique_vals) <= max_show:
        return ", ".join(unique_vals)
    return ", ".join(unique_vals[:max_show]) + f" +{len(unique_vals)-max_show}"


def read_workbook(path: Path) -> Tuple[Dict[str, int], List[FkEdge], List[InferredEdge]]:
    wb = load_workbook(path, read_only=True, data_only=True)

    ws_tables = wb["Tablolar"]
    ws_fk = wb["Iliskiler_FK"]
    ws_inf = wb["Iliskiler_Tahmini"]

    internal_tables: Dict[str, int] = {}
    for row in ws_tables.iter_rows(min_row=2, values_only=True):
        schema, table, col_count = row[0], row[1], row[2]
        if not schema or not table:
            continue
        internal_tables[f"{schema}.{table}"] = int(col_count or 0)

    fk_edges: List[FkEdge] = []
    for row in ws_fk.iter_rows(min_row=2, values_only=True):
        constraint, src_schema, src_table, src_cols, tgt_schema, tgt_table, _, target_in_script, _ = row
        if not (src_schema and src_table and tgt_schema and tgt_table):
            continue
        fk_edges.append(
            (
                f"{src_schema}.{src_table}",
                f"{tgt_schema}.{tgt_table}",
                str(src_cols or ""),
                str(target_in_script or "").strip().lower() == "yes",
                str(constraint or ""),
            )
        )

    inferred_edges: List[InferredEdge] = []
    for row in ws_inf.iter_rows(min_row=2, values_only=True):
        src, src_col, tgt, confidence, _ = row
        if not (src and src_col and tgt):
            continue
        inferred_edges.append((str(src), str(tgt), str(src_col), str(confidence or "")))

    return internal_tables, fk_edges, inferred_edges


def circle_positions(nodes: Sequence[str], center_x: float, center_y: float, radius: float) -> Dict[str, Tuple[float, float]]:
    pos: Dict[str, Tuple[float, float]] = {}
    if not nodes:
        return pos
    count = len(nodes)
    for i, node in enumerate(nodes):
        angle = (2.0 * 3.141592653589793 * i) / count
        pos[node] = (center_x + radius * float(__import__("math").cos(angle)), center_y + radius * float(__import__("math").sin(angle)))
    return pos


def merged_graph_for_fk(
    internal_tables: Dict[str, int],
    fk_edges: Sequence[FkEdge],
) -> Tuple[nx.DiGraph, Dict[Tuple[str, str], str]]:
    g = nx.DiGraph()
    for full_name, col_count in internal_tables.items():
        g.add_node(full_name, internal=True, col_count=col_count)

    edge_cols = defaultdict(list)
    edge_constraints = defaultdict(list)
    for src, dst, src_cols, target_in_script, constraint in fk_edges:
        if dst not in g.nodes:
            g.add_node(dst, internal=target_in_script, col_count=0)
        if src not in g.nodes:
            g.add_node(src, internal=True, col_count=0)
        key = (src, dst)
        edge_cols[key].append(src_cols)
        edge_constraints[key].append(constraint)

    for (src, dst), cols in edge_cols.items():
        g.add_edge(src, dst, kind="fk", constraints=edge_constraints[(src, dst)])

    labels = {(src, dst): compress_columns(cols) for (src, dst), cols in edge_cols.items()}
    return g, labels


def add_inferred_edges(
    base_graph: nx.DiGraph,
    internal_tables: Dict[str, int],
    inferred_edges: Sequence[InferredEdge],
) -> Dict[Tuple[str, str], str]:
    inferred_map = defaultdict(list)
    for src, tgt, src_col, confidence in inferred_edges:
        if src == tgt:
            continue
        if src not in internal_tables or tgt not in internal_tables:
            continue
        inferred_map[(src, tgt)].append(f"{src_col} ({confidence[:1].upper()})")

    for (src, tgt), cols in inferred_map.items():
        if base_graph.has_edge(src, tgt):
            # Explicit FK already exists, keep it authoritative.
            continue
        base_graph.add_edge(src, tgt, kind="inferred")

    return {(src, tgt): compress_columns(cols, max_show=2) for (src, tgt), cols in inferred_map.items()}


def build_positions(graph: nx.DiGraph, internal_tables: Dict[str, int]) -> Dict[str, Tuple[float, float]]:
    internal_nodes = sorted([n for n in graph.nodes if n in internal_tables])
    external_nodes = sorted([n for n in graph.nodes if n not in internal_tables])

    pos = {}
    pos.update(circle_positions(internal_nodes, center_x=-1.8, center_y=0.0, radius=2.4))
    pos.update(circle_positions(external_nodes, center_x=3.8, center_y=0.0, radius=2.0))
    return pos


def draw_diagram(
    graph: nx.DiGraph,
    internal_tables: Dict[str, int],
    explicit_labels: Dict[Tuple[str, str], str],
    inferred_labels: Dict[Tuple[str, str], str],
    title: str,
    output_base: Path,
) -> None:
    pos = build_positions(graph, internal_tables)
    fig, ax = plt.subplots(figsize=(22, 14))
    ax.set_axis_off()

    internal_nodes = [n for n in graph.nodes if n in internal_tables]
    external_nodes = [n for n in graph.nodes if n not in internal_tables]

    internal_sizes = [2200 + internal_tables[n] * 2 for n in internal_nodes]
    external_sizes = [1900 for _ in external_nodes]

    nx.draw_networkx_nodes(
        graph,
        pos=pos,
        nodelist=internal_nodes,
        node_color="#4C78A8",
        node_size=internal_sizes,
        edgecolors="#1F2A35",
        linewidths=1.1,
        ax=ax,
    )
    nx.draw_networkx_nodes(
        graph,
        pos=pos,
        nodelist=external_nodes,
        node_color="#B8C1CC",
        node_size=external_sizes,
        edgecolors="#4A5560",
        linewidths=1.0,
        ax=ax,
    )

    labels = {n: short_name(n) for n in graph.nodes}
    nx.draw_networkx_labels(graph, pos=pos, labels=labels, font_size=8.8, font_color="#111111", ax=ax)

    fk_edges = [(u, v) for u, v, d in graph.edges(data=True) if d.get("kind") == "fk"]
    inferred_edges = [(u, v) for u, v, d in graph.edges(data=True) if d.get("kind") == "inferred"]

    if fk_edges:
        nx.draw_networkx_edges(
            graph,
            pos=pos,
            edgelist=fk_edges,
            arrows=True,
            arrowstyle="-|>",
            arrowsize=16,
            width=1.8,
            alpha=0.75,
            edge_color="#2F3A45",
            connectionstyle="arc3,rad=0.08",
            ax=ax,
        )
        fk_label_map = {(u, v): explicit_labels.get((u, v), "") for (u, v) in fk_edges if explicit_labels.get((u, v))}
        if fk_label_map:
            nx.draw_networkx_edge_labels(
                graph,
                pos=pos,
                edge_labels=fk_label_map,
                font_size=7.5,
                rotate=False,
                bbox={"facecolor": "white", "edgecolor": "none", "alpha": 0.8, "pad": 0.15},
                ax=ax,
            )

    if inferred_edges:
        nx.draw_networkx_edges(
            graph,
            pos=pos,
            edgelist=inferred_edges,
            arrows=True,
            arrowstyle="-|>",
            arrowsize=15,
            width=1.6,
            alpha=0.85,
            edge_color="#D97706",
            style="dashed",
            connectionstyle="arc3,rad=-0.14",
            ax=ax,
        )
        inf_label_map = {(u, v): inferred_labels.get((u, v), "") for (u, v) in inferred_edges if inferred_labels.get((u, v))}
        if inf_label_map:
            nx.draw_networkx_edge_labels(
                graph,
                pos=pos,
                edge_labels=inf_label_map,
                font_size=7.2,
                rotate=False,
                bbox={"facecolor": "#FFF7ED", "edgecolor": "none", "alpha": 0.82, "pad": 0.15},
                ax=ax,
            )

    fk_count = len(fk_edges)
    inferred_count = len(inferred_edges)
    ax.set_title(
        f"{title}\nInternal tablo: {len(internal_nodes)} | External referans tablo: {len(external_nodes)} | FK: {fk_count} | Tahmini: {inferred_count}",
        fontsize=14,
        pad=14,
    )

    # Legend
    legend_items = [
        plt.Line2D([0], [0], marker="o", color="w", label="Internal table", markerfacecolor="#4C78A8", markeredgecolor="#1F2A35", markersize=8),
        plt.Line2D([0], [0], marker="o", color="w", label="External reference table", markerfacecolor="#B8C1CC", markeredgecolor="#4A5560", markersize=8),
        plt.Line2D([0], [0], color="#2F3A45", lw=2, label="FK (explicit)"),
    ]
    if inferred_edges:
        legend_items.append(plt.Line2D([0], [0], color="#D97706", lw=2, linestyle="--", label="Inferred internal relationship"))

    ax.legend(handles=legend_items, loc="upper right", framealpha=0.95, fontsize=9)

    output_base.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output_base.with_suffix(".png"), dpi=420, format="png", bbox_inches="tight", pad_inches=0.2)
    fig.savefig(output_base.with_suffix(".svg"), format="svg", bbox_inches="tight", pad_inches=0.2)
    plt.close(fig)


def write_readme(
    output_dir: Path,
    workbook_path: Path,
    internal_count: int,
    fk_count: int,
    inferred_all_count: int,
    inferred_drawn_count: int,
) -> None:
    readme = output_dir / "README.md"
    lines = [
        "# customers-vs-core ER Diyagramlari",
        "",
        f"- Kaynak workbook: `{workbook_path}`",
        f"- Internal tablo sayisi: **{internal_count}**",
        f"- Explicit FK sayisi: **{fk_count}**",
        f"- Tahmini iliski sayisi (tum): **{inferred_all_count}**",
        f"- Diyagrama cizilen tahmini iliski: **{inferred_drawn_count}** (ic tablo + self olmayan)",
        "",
        "## Dosyalar",
        "",
        "- `customers-vs-core-er-fk-highres.png/.svg`: Sadece explicit FK iliskileri",
        "- `customers-vs-core-er-fk-plus-tahmini-highres.png/.svg`: FK + ic tahmini iliskiler",
        "",
        "Not: Turuncu kesikli cizgiler isim eslesmesiyle cikarilan tahmini iliskilerdir.",
        "",
    ]
    readme.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate ER diagrams for customers-vs-core")
    parser.add_argument(
        "--input-xlsx",
        default="docs/customers-vs-core-veri-sozlugu.xlsx",
        help="Input workbook path",
    )
    parser.add_argument(
        "--output-dir",
        default="docs/customers-vs-core-er",
        help="Output directory",
    )
    args = parser.parse_args()

    workbook_path = Path(args.input_xlsx)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    internal_tables, fk_edges, inferred_edges = read_workbook(workbook_path)

    fk_graph, fk_labels = merged_graph_for_fk(internal_tables, fk_edges)
    draw_diagram(
        graph=fk_graph,
        internal_tables=internal_tables,
        explicit_labels=fk_labels,
        inferred_labels={},
        title="customers-vs-core ER (Explicit FK)",
        output_base=output_dir / "customers-vs-core-er-fk-highres",
    )

    combined_graph = fk_graph.copy()
    inferred_labels = add_inferred_edges(combined_graph, internal_tables, inferred_edges)
    draw_diagram(
        graph=combined_graph,
        internal_tables=internal_tables,
        explicit_labels=fk_labels,
        inferred_labels=inferred_labels,
        title="customers-vs-core ER (FK + Tahmini Ic Iliskiler)",
        output_base=output_dir / "customers-vs-core-er-fk-plus-tahmini-highres",
    )

    drawn_inferred_count = len([(u, v) for u, v, d in combined_graph.edges(data=True) if d.get("kind") == "inferred"])
    write_readme(
        output_dir=output_dir,
        workbook_path=workbook_path,
        internal_count=len(internal_tables),
        fk_count=len(fk_edges),
        inferred_all_count=len(inferred_edges),
        inferred_drawn_count=drawn_inferred_count,
    )

    print(f"Input: {workbook_path}")
    print(f"Output dir: {output_dir}")
    print(f"Internal tables: {len(internal_tables)}")
    print(f"Explicit FK rows: {len(fk_edges)}")
    print(f"Inferred rows: {len(inferred_edges)}")
    print(f"Drawn inferred edges: {drawn_inferred_count}")


if __name__ == "__main__":
    main()
