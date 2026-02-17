#!/usr/bin/env python3
"""Generate CRM full data dictionary workbook from SQL Server script."""

from __future__ import annotations

import argparse
import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


@dataclass
class ColumnInfo:
    name: str
    data_type: str
    nullable: str
    is_computed: bool = False
    computed_expression: str = ""
    is_masked: bool = False
    is_identity: bool = False
    default_expression: str = ""
    default_constraint: str = ""


@dataclass
class UniqueConstraint:
    name: str
    columns: List[str] = field(default_factory=list)


@dataclass
class ForeignKeyInfo:
    constraint: str
    source_schema: str
    source_table: str
    source_columns: List[str]
    target_schema: str
    target_table: str
    target_columns: List[str]
    on_delete: str = ""
    on_update: str = ""

    @property
    def source_full(self) -> str:
        return f"{self.source_schema}.{self.source_table}"

    @property
    def target_full(self) -> str:
        return f"{self.target_schema}.{self.target_table}"


@dataclass
class TableInfo:
    schema: str
    table: str
    columns: List[ColumnInfo] = field(default_factory=list)
    pk_constraint: str = ""
    pk_columns: List[str] = field(default_factory=list)
    unique_constraints: List[UniqueConstraint] = field(default_factory=list)

    @property
    def full_name(self) -> str:
        return f"{self.schema}.{self.table}"


def split_batches(sql_text: str) -> List[str]:
    return [x.strip() for x in re.split(r"^\s*GO\s*$", sql_text, flags=re.IGNORECASE | re.MULTILINE) if x.strip()]


def extract_bracket_names(segment: str) -> List[str]:
    return [x.strip() for x in re.findall(r"\[([^\]]+)\]", segment, flags=re.IGNORECASE)]


def parse_parenthesized(text: str, start_idx: int) -> Tuple[str, int]:
    """Read balanced parenthesized expression from start_idx ('(' expected)."""
    if start_idx >= len(text) or text[start_idx] != "(":
        return "", start_idx
    i = start_idx
    depth = 0
    inside_quote = False
    while i < len(text):
        ch = text[i]
        if ch == "'" and (i + 1 >= len(text) or text[i + 1] != "'"):
            inside_quote = not inside_quote
        elif ch == "'" and i + 1 < len(text) and text[i + 1] == "'":
            i += 1
        elif not inside_quote:
            if ch == "(":
                depth += 1
            elif ch == ")":
                depth -= 1
                if depth == 0:
                    return text[start_idx : i + 1], i + 1
        i += 1
    return text[start_idx:], len(text)


def split_top_level_csv(text: str) -> List[str]:
    parts: List[str] = []
    current: List[str] = []
    depth = 0
    inside_quote = False
    i = 0
    while i < len(text):
        ch = text[i]
        if ch == "'" and (i + 1 >= len(text) or text[i + 1] != "'"):
            inside_quote = not inside_quote
        elif ch == "'" and i + 1 < len(text) and text[i + 1] == "'":
            current.append(ch)
            i += 1
            ch = text[i]
        if not inside_quote:
            if ch == "(":
                depth += 1
            elif ch == ")" and depth > 0:
                depth -= 1
            elif ch == "," and depth == 0:
                item = "".join(current).strip()
                if item:
                    parts.append(item)
                current = []
                i += 1
                continue
        current.append(ch)
        i += 1
    tail = "".join(current).strip()
    if tail:
        parts.append(tail)
    return parts


def find_create_table_blocks(batch: str) -> List[Tuple[str, str, str]]:
    blocks: List[Tuple[str, str, str]] = []
    pattern = re.compile(
        r"CREATE\s+TABLE\s+\[(?P<schema>[^\]]+)\]\.\[(?P<table>[^\]]+)\]\s*\(",
        flags=re.IGNORECASE,
    )
    for m in pattern.finditer(batch):
        schema = m.group("schema").strip()
        table = m.group("table").strip()
        start = m.end() - 1  # at "("
        body_with_paren, _ = parse_parenthesized(batch, start)
        if not body_with_paren:
            continue
        body = body_with_paren[1:-1]  # remove outer ()
        blocks.append((schema, table, body))
    return blocks


def parse_inline_default(rest: str) -> Tuple[str, str]:
    upper = rest.upper()
    idx = upper.find(" DEFAULT ")
    if idx < 0:
        if upper.startswith("DEFAULT "):
            idx = 0
        else:
            return "", ""

    default_part = rest[idx + (1 if idx > 0 else 0) :]
    m_constraint = re.search(r"CONSTRAINT\s+\[([^\]]+)\]\s+DEFAULT", rest, flags=re.IGNORECASE)
    constraint_name = m_constraint.group(1).strip() if m_constraint else ""

    m_default = re.search(r"\bDEFAULT\b", default_part, flags=re.IGNORECASE)
    if not m_default:
        return "", constraint_name
    expr_part = default_part[m_default.end() :].strip()

    if expr_part.startswith("("):
        expr, _ = parse_parenthesized(expr_part, 0)
        return expr.strip(), constraint_name

    token = expr_part.split()[0] if expr_part else ""
    return token.strip(), constraint_name


def parse_column(entry: str) -> ColumnInfo | None:
    line = " ".join(entry.replace("\r", " ").replace("\n", " ").split())

    m_computed = re.match(r"^\[(?P<name>[^\]]+)\]\s+AS\s+(?P<expr>.+)$", line, flags=re.IGNORECASE)
    if m_computed:
        return ColumnInfo(
            name=m_computed.group("name").strip(),
            data_type="COMPUTED",
            nullable="N/A",
            is_computed=True,
            computed_expression=m_computed.group("expr").strip(),
        )

    m_col = re.match(r"^\[(?P<name>[^\]]+)\]\s+\[(?P<dtype>[^\]]+)\](?P<rest>.*)$", line, flags=re.IGNORECASE)
    if not m_col:
        return None

    name = m_col.group("name").strip()
    dtype = m_col.group("dtype").strip()
    rest = m_col.group("rest").strip()

    size_part = ""
    if rest.startswith("("):
        size_part, consumed = parse_parenthesized(rest, 0)
        rest = rest[consumed:].strip()

    data_type = f"{dtype}{size_part}"
    nullable = "UNKNOWN"
    if re.search(r"\bNOT\s+NULL\b", rest, flags=re.IGNORECASE):
        nullable = "NOT NULL"
    elif re.search(r"\bNULL\b", rest, flags=re.IGNORECASE):
        nullable = "NULL"

    default_expr, default_constraint = parse_inline_default(rest)

    return ColumnInfo(
        name=name,
        data_type=data_type,
        nullable=nullable,
        is_masked=bool(re.search(r"\bMASKED\s+WITH\b", rest, flags=re.IGNORECASE)),
        is_identity=bool(re.search(r"\bIDENTITY\s*\(", rest, flags=re.IGNORECASE)),
        default_expression=default_expr,
        default_constraint=default_constraint,
    )


def parse_create_table_entry_constraints(entry: str, table: TableInfo) -> None:
    text = entry.strip()
    if not text.upper().startswith("CONSTRAINT"):
        return

    m_pk = re.search(
        r"CONSTRAINT\s+\[(?P<name>[^\]]+)\]\s+PRIMARY\s+KEY\b[^(]*\((?P<cols>.*?)\)",
        text,
        flags=re.IGNORECASE | re.DOTALL,
    )
    if m_pk:
        table.pk_constraint = m_pk.group("name").strip()
        table.pk_columns = extract_bracket_names(m_pk.group("cols"))
        return

    m_uq = re.search(
        r"CONSTRAINT\s+\[(?P<name>[^\]]+)\]\s+UNIQUE\b[^(]*\((?P<cols>.*?)\)",
        text,
        flags=re.IGNORECASE | re.DOTALL,
    )
    if m_uq:
        table.unique_constraints.append(
            UniqueConstraint(
                name=m_uq.group("name").strip(),
                columns=extract_bracket_names(m_uq.group("cols")),
            )
        )


def parse_tables_from_batches(batches: Sequence[str]) -> Dict[str, TableInfo]:
    tables: Dict[str, TableInfo] = {}
    for batch in batches:
        for schema, table_name, body in find_create_table_blocks(batch):
            info = TableInfo(schema=schema, table=table_name)
            for entry in split_top_level_csv(body):
                entry_stripped = entry.strip()
                if entry_stripped.upper().startswith("CONSTRAINT"):
                    parse_create_table_entry_constraints(entry_stripped, info)
                    continue
                col = parse_column(entry_stripped)
                if col:
                    info.columns.append(col)
            tables[info.full_name] = info
    return dict(sorted(tables.items(), key=lambda kv: kv[0].lower()))


def parse_alter_defaults(batch: str) -> Tuple[str, str, str, str]:
    m_hdr = re.search(r"ALTER\s+TABLE\s+\[(?P<schema>[^\]]+)\]\.\[(?P<table>[^\]]+)\]", batch, flags=re.IGNORECASE)
    if not m_hdr:
        return "", "", "", ""
    m_def = re.search(
        r"(?:CONSTRAINT\s+\[(?P<constraint>[^\]]+)\]\s+)?DEFAULT\s+(?P<expr>.+?)\s+FOR\s+\[(?P<column>[^\]]+)\]",
        batch,
        flags=re.IGNORECASE | re.DOTALL,
    )
    if not m_def:
        return "", "", "", ""
    full_table = f"{m_hdr.group('schema').strip()}.{m_hdr.group('table').strip()}"
    column = m_def.group("column").strip()
    expr = " ".join((m_def.group("expr") or "").split())
    constraint = (m_def.group("constraint") or "").strip()
    return full_table, column, expr, constraint


def parse_alter_pk(batch: str) -> Tuple[str, str, List[str]]:
    m_hdr = re.search(r"ALTER\s+TABLE\s+\[(?P<schema>[^\]]+)\]\.\[(?P<table>[^\]]+)\]", batch, flags=re.IGNORECASE)
    if not m_hdr:
        return "", "", []
    m_pk = re.search(
        r"CONSTRAINT\s+\[(?P<constraint>[^\]]+)\]\s+PRIMARY\s+KEY\b[^(]*\((?P<cols>.*?)\)",
        batch,
        flags=re.IGNORECASE | re.DOTALL,
    )
    if not m_pk:
        return "", "", []
    full_table = f"{m_hdr.group('schema').strip()}.{m_hdr.group('table').strip()}"
    return full_table, m_pk.group("constraint").strip(), extract_bracket_names(m_pk.group("cols"))


def parse_alter_unique(batch: str) -> Tuple[str, str, List[str]]:
    m_hdr = re.search(r"ALTER\s+TABLE\s+\[(?P<schema>[^\]]+)\]\.\[(?P<table>[^\]]+)\]", batch, flags=re.IGNORECASE)
    if not m_hdr:
        return "", "", []
    m_uq = re.search(
        r"CONSTRAINT\s+\[(?P<constraint>[^\]]+)\]\s+UNIQUE\b[^(]*\((?P<cols>.*?)\)",
        batch,
        flags=re.IGNORECASE | re.DOTALL,
    )
    if not m_uq:
        return "", "", []
    full_table = f"{m_hdr.group('schema').strip()}.{m_hdr.group('table').strip()}"
    return full_table, m_uq.group("constraint").strip(), extract_bracket_names(m_uq.group("cols"))


def parse_alter_fk(batch: str) -> ForeignKeyInfo | None:
    m_hdr = re.search(
        r"ALTER\s+TABLE\s+\[(?P<src_schema>[^\]]+)\]\.\[(?P<src_table>[^\]]+)\]",
        batch,
        flags=re.IGNORECASE,
    )
    if not m_hdr:
        return None

    m_fk = re.search(
        r"CONSTRAINT\s+\[(?P<constraint>[^\]]+)\]\s+FOREIGN\s+KEY\s*\((?P<src_cols>.*?)\)\s*REFERENCES\s+\[(?P<ref_schema>[^\]]+)\]\.\[(?P<ref_table>[^\]]+)\]\s*\((?P<ref_cols>.*?)\)",
        batch,
        flags=re.IGNORECASE | re.DOTALL,
    )
    if not m_fk:
        return None

    upper = batch.upper()
    m_delete = re.search(
        r"ON\s+DELETE\s+(CASCADE|SET\s+NULL|SET\s+DEFAULT|NO\s+ACTION)",
        upper,
        flags=re.IGNORECASE,
    )
    m_update = re.search(
        r"ON\s+UPDATE\s+(CASCADE|SET\s+NULL|SET\s+DEFAULT|NO\s+ACTION)",
        upper,
        flags=re.IGNORECASE,
    )

    return ForeignKeyInfo(
        constraint=m_fk.group("constraint").strip(),
        source_schema=m_hdr.group("src_schema").strip(),
        source_table=m_hdr.group("src_table").strip(),
        source_columns=extract_bracket_names(m_fk.group("src_cols")),
        target_schema=m_fk.group("ref_schema").strip(),
        target_table=m_fk.group("ref_table").strip(),
        target_columns=extract_bracket_names(m_fk.group("ref_cols")),
        on_delete=(m_delete.group(1).strip().upper() if m_delete else ""),
        on_update=(m_update.group(1).strip().upper() if m_update else ""),
    )


def enrich_tables_from_alter_batches(
    batches: Sequence[str],
    tables: Dict[str, TableInfo],
) -> Tuple[Dict[Tuple[str, str], Tuple[str, str]], List[ForeignKeyInfo]]:
    defaults: Dict[Tuple[str, str], Tuple[str, str]] = {}
    fks: List[ForeignKeyInfo] = []

    for batch in batches:
        upper = batch.upper()
        if not upper.startswith("ALTER TABLE"):
            continue

        if "FOREIGN KEY" in upper and "REFERENCES" in upper:
            fk = parse_alter_fk(batch)
            if fk:
                fks.append(fk)
            continue

        if "DEFAULT" in upper and "FOR [" in upper:
            full_table, column, expr, constraint = parse_alter_defaults(batch)
            if full_table and column:
                defaults[(full_table.upper(), column.upper())] = (expr, constraint)
            continue

        if "PRIMARY KEY" in upper:
            full_table, constraint, cols = parse_alter_pk(batch)
            if full_table in tables and constraint:
                tables[full_table].pk_constraint = constraint
                tables[full_table].pk_columns = cols
            continue

        if "UNIQUE" in upper and "CONSTRAINT" in upper:
            full_table, constraint, cols = parse_alter_unique(batch)
            if full_table in tables and constraint:
                exists = any(uq.name == constraint for uq in tables[full_table].unique_constraints)
                if not exists:
                    tables[full_table].unique_constraints.append(UniqueConstraint(name=constraint, columns=cols))

    fks.sort(key=lambda x: (x.source_schema, x.source_table, x.constraint))
    return defaults, fks


def style_sheet(ws, widths: Sequence[int]) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    center = Alignment(horizontal="center", vertical="center")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def unique_column_set(unique_constraints: Sequence[UniqueConstraint]) -> set[str]:
    cols = set()
    for uq in unique_constraints:
        for col in uq.columns:
            cols.add(col.upper())
    return cols


def build_workbook(
    source_file: Path,
    output_file: Path,
    tables: Dict[str, TableInfo],
    alter_defaults: Dict[Tuple[str, str], Tuple[str, str]],
    foreign_keys: Sequence[ForeignKeyInfo],
) -> None:
    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "Ozet"
    ws_tables = wb.create_sheet("Tablolar")
    ws_columns = wb.create_sheet("Kolonlar")
    ws_fk = wb.create_sheet("Iliskiler_FK")
    ws_fk_summary = wb.create_sheet("Iliskiler_Ozet")

    fk_out = defaultdict(int)
    fk_in = defaultdict(int)
    for fk in foreign_keys:
        fk_out[fk.source_full] += 1
        fk_in[fk.target_full] += 1

    default_count = 0
    total_columns = sum(len(t.columns) for t in tables.values())
    total_uniques = sum(len(t.unique_constraints) for t in tables.values())
    for t in tables.values():
        for c in t.columns:
            has_inline = bool(c.default_expression)
            has_alter = (t.full_name.upper(), c.name.upper()) in alter_defaults
            if has_inline or has_alter:
                default_count += 1

    ws_summary.append(["Alan", "Deger"])
    ws_summary.append(["Kaynak Dosya", str(source_file)])
    ws_summary.append(["Uretim Zamani", datetime.now().isoformat(timespec="seconds")])
    ws_summary.append(["Toplam Tablo", len(tables)])
    ws_summary.append(["Toplam Kolon", total_columns])
    ws_summary.append(["Toplam PK", sum(1 for t in tables.values() if t.pk_columns)])
    ws_summary.append(["Toplam Unique Constraint", total_uniques])
    ws_summary.append(["Default Tanimli Kolon", default_count])
    ws_summary.append(["Toplam FK (Explicit)", len(foreign_keys)])
    ws_summary.append(["FK Kaynagi Olan Tablo", len({fk.source_full for fk in foreign_keys})])
    ws_summary.append(["FK Hedefi Olan Tablo", len({fk.target_full for fk in foreign_keys})])
    style_sheet(ws_summary, [34, 95])

    ws_tables.append(
        [
            "Schema",
            "Table",
            "ColumnCount",
            "PKConstraint",
            "PKColumns",
            "UniqueConstraintCount",
            "DefaultColumnCount",
            "FKOutCount",
            "FKInCount",
        ]
    )
    for full, table in tables.items():
        default_columns = 0
        for col in table.columns:
            has_inline = bool(col.default_expression)
            has_alter = (full.upper(), col.name.upper()) in alter_defaults
            if has_inline or has_alter:
                default_columns += 1
        ws_tables.append(
            [
                table.schema,
                table.table,
                len(table.columns),
                table.pk_constraint,
                ", ".join(table.pk_columns),
                len(table.unique_constraints),
                default_columns,
                fk_out.get(full, 0),
                fk_in.get(full, 0),
            ]
        )
    style_sheet(ws_tables, [14, 33, 12, 34, 30, 23, 18, 12, 12])

    ws_columns.append(
        [
            "Schema",
            "Table",
            "Ordinal",
            "Column",
            "DataType",
            "Nullability",
            "IsPrimaryKey",
            "IsUniqueMember",
            "IsComputed",
            "ComputedExpression",
            "IsMasked",
            "IsIdentity",
            "DefaultExpression",
            "DefaultConstraint",
        ]
    )

    for full, table in tables.items():
        pk_set = {x.upper() for x in table.pk_columns}
        uq_set = unique_column_set(table.unique_constraints)
        for idx, col in enumerate(table.columns, start=1):
            default_expr = col.default_expression
            default_constraint = col.default_constraint
            if (full.upper(), col.name.upper()) in alter_defaults:
                d_expr, d_const = alter_defaults[(full.upper(), col.name.upper())]
                default_expr = d_expr or default_expr
                default_constraint = d_const or default_constraint

            ws_columns.append(
                [
                    table.schema,
                    table.table,
                    idx,
                    col.name,
                    col.data_type,
                    col.nullable,
                    "Yes" if col.name.upper() in pk_set else "No",
                    "Yes" if col.name.upper() in uq_set else "No",
                    "Yes" if col.is_computed else "No",
                    col.computed_expression,
                    "Yes" if col.is_masked else "No",
                    "Yes" if col.is_identity else "No",
                    default_expr,
                    default_constraint,
                ]
            )
    style_sheet(ws_columns, [12, 30, 10, 36, 24, 12, 12, 13, 11, 65, 10, 10, 35, 36])

    ws_fk.append(
        [
            "Constraint",
            "SourceSchema",
            "SourceTable",
            "SourceColumns",
            "TargetSchema",
            "TargetTable",
            "TargetColumns",
            "DeleteAction",
            "UpdateAction",
            "TargetInScript",
        ]
    )
    table_keys_upper = {x.upper() for x in tables.keys()}
    for fk in foreign_keys:
        ws_fk.append(
            [
                fk.constraint,
                fk.source_schema,
                fk.source_table,
                ", ".join(fk.source_columns),
                fk.target_schema,
                fk.target_table,
                ", ".join(fk.target_columns),
                fk.on_delete,
                fk.on_update,
                "Yes" if fk.target_full.upper() in table_keys_upper else "No",
            ]
        )
    style_sheet(ws_fk, [48, 14, 30, 30, 14, 30, 30, 16, 16, 14])

    ws_fk_summary.append(
        [
            "SourceTable",
            "TargetTable",
            "FKCount",
            "Constraints",
            "SourceColumns",
        ]
    )
    pair_map: Dict[Tuple[str, str], List[ForeignKeyInfo]] = defaultdict(list)
    for fk in foreign_keys:
        pair_map[(fk.source_full, fk.target_full)].append(fk)
    for (source, target), items in sorted(pair_map.items(), key=lambda kv: (kv[0][0], kv[0][1])):
        ws_fk_summary.append(
            [
                source,
                target,
                len(items),
                "; ".join(i.constraint for i in items),
                "; ".join(", ".join(i.source_columns) for i in items),
            ]
        )
    style_sheet(ws_fk_summary, [38, 38, 10, 75, 75])

    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_file)


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate full CRM data dictionary workbook")
    parser.add_argument(
        "--input",
        default="/home/ubuntu/.cursor/projects/workspace/uploads/script_2.17.26_generatescript.txt",
        help="Input SQL script file",
    )
    parser.add_argument(
        "--output",
        default="docs/crm-full-veri-sozlugu.xlsx",
        help="Output xlsx file path",
    )
    args = parser.parse_args()

    input_file = Path(args.input)
    output_file = Path(args.output)
    if not input_file.exists():
        raise FileNotFoundError(f"Input file not found: {input_file}")

    sql_text = input_file.read_text(encoding="utf-8", errors="ignore")
    batches = split_batches(sql_text)
    tables = parse_tables_from_batches(batches)
    alter_defaults, foreign_keys = enrich_tables_from_alter_batches(batches, tables)

    build_workbook(
        source_file=input_file,
        output_file=output_file,
        tables=tables,
        alter_defaults=alter_defaults,
        foreign_keys=foreign_keys,
    )

    print(f"Input: {input_file}")
    print(f"Output: {output_file}")
    print(f"Batches: {len(batches)}")
    print(f"Tables: {len(tables)}")
    print(f"Columns: {sum(len(t.columns) for t in tables.values())}")
    print(f"Defaults (ALTER): {len(alter_defaults)}")
    print(f"Foreign Keys: {len(foreign_keys)}")


if __name__ == "__main__":
    main()
