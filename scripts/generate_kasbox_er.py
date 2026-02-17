#!/usr/bin/env python3
"""Generate readable, high-resolution ER diagram assets for Kasbox."""

from __future__ import annotations

import argparse
import csv
import math
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Sequence, Set, Tuple

import matplotlib.pyplot as plt
import networkx as nx
from openpyxl import load_workbook

Relationship = Tuple[str, str, str, str]  # (source, target, via_column, confidence)

CATEGORY_COLORS = {
    "IBS": "#4C78A8",
    "KBAUDIT": "#F58518",
    "KBARSIV": "#72B7B2",
    "KBFAST": "#54A24B",
    "KBSIPER": "#E45756",
    "KAS_MSG": "#B279A2",
    "TCMB": "#FF9DA6",
    "KB_GENEL": "#9D755D",
    "DIGER": "#BAB0AC",
}

HUB_COLUMNS = ("STATU", "SERVISKODU", "DPK")


@dataclass
class TableInfo:
    schema: str
    table: str
    columns: Set[str]
    pk_columns: List[str]

    @property
    def full_name(self) -> str:
        return f"{self.schema}.{self.table}"


def short_name(full_table_name: str) -> str:
    return full_table_name.split(".", 1)[1]


def table_category(full_table_name: str) -> str:
    table = short_name(full_table_name)
    if table.startswith("IBS_"):
        return "IBS"
    if table.startswith("KBAUDIT_"):
        return "KBAUDIT"
    if table.startswith("KBARSIV_"):
        return "KBARSIV"
    if table.startswith("KBFAST"):
        return "KBFAST"
    if table.startswith("KBSIPER"):
        return "KBSIPER"
    if table.startswith("ALL_MSG_") or table.startswith("FROM_") or table.startswith("TO_"):
        return "KAS_MSG"
    if table.startswith("TCMB"):
        return "TCMB"
    if table.startswith("KB"):
        return "KB_GENEL"
    return "DIGER"


def read_dictionary(input_xlsx: Path) -> Dict[str, TableInfo]:
    wb = load_workbook(input_xlsx, data_only=True, read_only=True)
    ws = wb["Kolonlar"]

    table_columns: Dict[str, Set[str]] = defaultdict(set)
    table_pks: Dict[str, List[str]] = defaultdict(list)
    table_parts: Dict[str, Tuple[str, str]] = {}

    for row in ws.iter_rows(min_row=2, values_only=True):
        schema, table, column, _, _, _, is_pk, _, _ = row
        if not (schema and table and column):
            continue
        full = f"{schema}.{table}"
        table_parts[full] = (str(schema), str(table))
        table_columns[full].add(str(column))
        if str(is_pk).strip().lower() == "yes":
            table_pks[full].append(str(column))

    table_infos: Dict[str, TableInfo] = {}
    for full, cols in table_columns.items():
        schema, table = table_parts[full]
        table_infos[full] = TableInfo(
            schema=schema,
            table=table,
            columns=cols,
            pk_columns=table_pks.get(full, []),
        )
    return table_infos


def infer_relationships(tables: Dict[str, TableInfo]) -> List[Relationship]:
    """Infer relationship candidates by single-PK name matching."""
    single_pk_index: Dict[str, List[str]] = defaultdict(list)
    for full, info in tables.items():
        if len(info.pk_columns) == 1:
            single_pk_index[info.pk_columns[0].upper()].append(full)

    relationships: List[Relationship] = []
    seen: Set[Tuple[str, str, str]] = set()

    for source_full, source_info in tables.items():
        source_pk_set = {x.upper() for x in source_info.pk_columns}
        for col in source_info.columns:
            col_upper = col.upper()
            if col_upper in source_pk_set:
                continue
            targets = [t for t in single_pk_index.get(col_upper, []) if t != source_full]
            if len(targets) != 1:
                continue

            target = targets[0]
            confidence = "high" if (col_upper.endswith("_ID") or col_upper.endswith("ID")) else "medium"
            key = (source_full, target, col)
            if key in seen:
                continue

            relationships.append((source_full, target, col, confidence))
            seen.add(key)

    relationships.sort(key=lambda x: (x[2], x[0], x[1]))
    return relationships


def split_relationships(relationships: Sequence[Relationship]) -> Dict[str, List[Relationship]]:
    buckets: Dict[str, List[Relationship]] = {key: [] for key in HUB_COLUMNS}
    buckets["OTHER"] = []
    for rel in relationships:
        via = rel[2].upper()
        if via in buckets:
            buckets[via].append(rel)
        else:
            buckets["OTHER"].append(rel)
    for key in buckets:
        buckets[key].sort(key=lambda x: (x[0], x[1], x[2]))
    return buckets


def write_relationship_csv(relationships: Sequence[Relationship], output_csv: Path) -> None:
    output_csv.parent.mkdir(parents=True, exist_ok=True)
    with output_csv.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["source_table", "target_table", "via_column", "confidence"])
        writer.writerows(relationships)


def save_figure(fig: plt.Figure, output_base: Path) -> Tuple[Path, Path]:
    png = output_base.with_suffix(".png")
    svg = output_base.with_suffix(".svg")
    output_base.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(png, format="png", dpi=420, bbox_inches="tight", pad_inches=0.25)
    fig.savefig(svg, format="svg", bbox_inches="tight", pad_inches=0.25)
    plt.close(fig)
    return png, svg


def node_color(full_table_name: str) -> str:
    return CATEGORY_COLORS.get(table_category(full_table_name), CATEGORY_COLORS["DIGER"])


def draw_category_summary(
    tables: Dict[str, TableInfo], relationships: Sequence[Relationship], output_base: Path
) -> Tuple[Path, Path]:
    table_counts = Counter(table_category(full) for full in tables)
    edge_counts = Counter((table_category(src), table_category(dst)) for src, dst, _, _ in relationships)

    graph = nx.DiGraph()
    for category, count in sorted(table_counts.items()):
        graph.add_node(category, table_count=count)
    for (src_cat, dst_cat), edge_count in sorted(edge_counts.items()):
        graph.add_edge(src_cat, dst_cat, weight=edge_count)

    pos = nx.circular_layout(graph, scale=4.0)
    fig, ax = plt.subplots(figsize=(18, 12))
    ax.set_axis_off()

    node_sizes = [2200 + graph.nodes[n]["table_count"] * 85 for n in graph.nodes]
    node_colors = [CATEGORY_COLORS.get(n, CATEGORY_COLORS["DIGER"]) for n in graph.nodes]

    nx.draw_networkx_nodes(
        graph,
        pos,
        node_size=node_sizes,
        node_color=node_colors,
        edgecolors="#1F2A35",
        linewidths=1.1,
        ax=ax,
    )

    widths = [0.8 + graph.edges[e]["weight"] * 0.45 for e in graph.edges]
    nx.draw_networkx_edges(
        graph,
        pos,
        arrows=True,
        arrowstyle="-|>",
        arrowsize=18,
        width=widths,
        alpha=0.55,
        edge_color="#3B4A5A",
        connectionstyle="arc3,rad=0.15",
        ax=ax,
    )

    labels = {n: f"{n}\n({graph.nodes[n]['table_count']} tablo)" for n in graph.nodes}
    nx.draw_networkx_labels(graph, pos, labels=labels, font_size=10, font_weight="bold", ax=ax)

    edge_labels = {(u, v): graph.edges[(u, v)]["weight"] for u, v in graph.edges}
    nx.draw_networkx_edge_labels(
        graph,
        pos,
        edge_labels=edge_labels,
        font_size=9,
        rotate=False,
        bbox={"facecolor": "white", "edgecolor": "none", "alpha": 0.72, "pad": 0.2},
        ax=ax,
    )

    ax.set_title(
        (
            "Kasbox ER - Kategori Ozeti\n"
            "Node: kategori ve tablo sayisi | Edge etiketi: kategori arasi iliski sayisi"
        ),
        fontsize=15,
        pad=15,
    )

    return save_figure(fig, output_base)


def draw_hub_star(
    relationships: Sequence[Relationship], hub_column: str, output_base: Path
) -> Tuple[Path, Path]:
    target_counts = Counter(dst for _, dst, _, _ in relationships)
    hub_target, _ = target_counts.most_common(1)[0]
    source_tables = sorted({src for src, dst, _, _ in relationships if dst == hub_target})

    fig, ax = plt.subplots(figsize=(20, 20))
    ax.set_axis_off()

    radius = max(7.0, 0.55 * len(source_tables))
    pos: Dict[str, Tuple[float, float]] = {hub_target: (0.0, 0.0)}
    for i, src in enumerate(source_tables):
        angle = (2.0 * math.pi * i) / max(1, len(source_tables))
        pos[src] = (radius * math.cos(angle), radius * math.sin(angle))

    graph = nx.DiGraph()
    graph.add_node(hub_target)
    for src in source_tables:
        graph.add_node(src)
    for src, dst, _, confidence in relationships:
        graph.add_edge(src, dst, confidence=confidence)

    node_sizes = []
    node_colors = []
    for n in graph.nodes:
        if n == hub_target:
            node_sizes.append(4100)
            node_colors.append("#2E3A46")
        else:
            node_sizes.append(1550)
            node_colors.append(node_color(n))

    nx.draw_networkx_nodes(
        graph,
        pos,
        node_size=node_sizes,
        node_color=node_colors,
        edgecolors="#1F2A35",
        linewidths=1.1,
        ax=ax,
    )
    nx.draw_networkx_edges(
        graph,
        pos,
        arrows=True,
        arrowstyle="-|>",
        arrowsize=16,
        width=1.2,
        alpha=0.55,
        edge_color="#3B4A5A",
        connectionstyle="arc3,rad=0.07",
        ax=ax,
    )

    labels = {n: short_name(n) for n in graph.nodes}
    nx.draw_networkx_labels(graph, pos, labels=labels, font_size=8.4, font_color="#111111", ax=ax)
    nx.draw_networkx_labels(
        graph,
        pos={hub_target: (0.0, 0.0)},
        labels={hub_target: short_name(hub_target)},
        font_size=10.2,
        font_color="#FFFFFF",
        font_weight="bold",
        ax=ax,
    )

    ax.set_title(
        (
            f"Kasbox ER - {hub_column} Hub Iliskileri\n"
            f"Merkez tablo: {hub_target} | Bagli tablo: {len(source_tables)} | Iliski: {len(relationships)}"
        ),
        fontsize=15,
        pad=16,
    )

    return save_figure(fig, output_base)


def draw_other_relationships(
    relationships: Sequence[Relationship], output_base: Path
) -> Tuple[Path, Path]:
    graph = nx.DiGraph()
    for src, dst, via, confidence in relationships:
        graph.add_node(src)
        graph.add_node(dst)
        graph.add_edge(src, dst, via=via, confidence=confidence)

    pos = nx.spring_layout(graph, seed=41, k=1.25, iterations=550)

    fig, ax = plt.subplots(figsize=(18, 12))
    ax.set_axis_off()

    nx.draw_networkx_nodes(
        graph,
        pos,
        node_size=1900,
        node_color=[node_color(n) for n in graph.nodes],
        edgecolors="#1F2A35",
        linewidths=1.0,
        ax=ax,
    )
    nx.draw_networkx_edges(
        graph,
        pos,
        arrows=True,
        arrowstyle="-|>",
        arrowsize=15,
        width=1.4,
        alpha=0.75,
        edge_color="#3B4A5A",
        connectionstyle="arc3,rad=0.11",
        ax=ax,
    )
    labels = {n: short_name(n) for n in graph.nodes}
    nx.draw_networkx_labels(graph, pos, labels=labels, font_size=9.5, ax=ax)

    edge_labels = {(u, v): d["via"] for u, v, d in graph.edges(data=True)}
    nx.draw_networkx_edge_labels(
        graph,
        pos,
        edge_labels=edge_labels,
        font_size=9,
        rotate=False,
        bbox={"facecolor": "white", "edgecolor": "none", "alpha": 0.82, "pad": 0.22},
        ax=ax,
    )

    ax.set_title(
        (
            "Kasbox ER - Diger Iliskiler\n"
            f"Toplam node: {graph.number_of_nodes()} | Toplam iliski: {graph.number_of_edges()}"
        ),
        fontsize=15,
        pad=15,
    )
    return save_figure(fig, output_base)


def write_readme(
    output_dir: Path,
    tables: Dict[str, TableInfo],
    relationships: Sequence[Relationship],
    buckets: Dict[str, List[Relationship]],
) -> Path:
    path = output_dir / "README.md"
    lines: List[str] = []
    lines.append("# Kasbox ER Semasi - Okunabilir Cikti Paketi")
    lines.append("")
    lines.append(
        "Not: Kaynak SQL scriptte acik FOREIGN KEY olmadigi icin iliskiler, "
        "tekil PK ile kolon isim eslesmesine gore otomatik cikartilmistir."
    )
    lines.append("")
    lines.append(f"- Toplam tablo: **{len(tables)}**")
    lines.append(f"- Toplam cikartilan iliski: **{len(relationships)}**")
    lines.append("")
    lines.append("## Dosyalar")
    lines.append("")
    lines.append("- `00-kategori-ozet-highres.png/.svg`")
    lines.append("- `01-statu-hub-highres.png/.svg`")
    lines.append("- `02-serviskodu-hub-highres.png/.svg`")
    lines.append("- `03-dpk-hub-highres.png/.svg`")
    lines.append("- `04-diger-iliskiler-highres.png/.svg`")
    lines.append("- `99-iliskiler-tam.csv`")
    lines.append("- `statu-iliskileri.csv`")
    lines.append("- `serviskodu-iliskileri.csv`")
    lines.append("- `dpk-iliskileri.csv`")
    lines.append("- `diger-iliskiler.csv`")
    lines.append("")
    lines.append("## Iliski Dagilimi")
    lines.append("")
    for key in ("STATU", "SERVISKODU", "DPK", "OTHER"):
        lines.append(f"- {key}: {len(buckets[key])}")
    lines.append("")
    lines.append("## Renk Anahtari")
    lines.append("")
    for cat in sorted(CATEGORY_COLORS):
        lines.append(f"- {cat}")
    lines.append("")

    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate readable Kasbox ER assets.")
    parser.add_argument(
        "--input-xlsx",
        default="docs/kasbox-veri-sozlugu.xlsx",
        help="Input data dictionary xlsx path",
    )
    parser.add_argument(
        "--output-dir",
        default="docs/kasbox-er-okunabilir",
        help="Output folder for ER assets",
    )
    args = parser.parse_args()

    input_xlsx = Path(args.input_xlsx)
    if not input_xlsx.exists():
        raise FileNotFoundError(f"Input dictionary not found: {input_xlsx}")

    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    tables = read_dictionary(input_xlsx)
    relationships = infer_relationships(tables)
    buckets = split_relationships(relationships)

    generated: List[Path] = []
    generated.extend(
        draw_category_summary(tables, relationships, output_dir / "00-kategori-ozet-highres")
    )
    generated.extend(
        draw_hub_star(buckets["STATU"], "STATU", output_dir / "01-statu-hub-highres")
    )
    generated.extend(
        draw_hub_star(
            buckets["SERVISKODU"], "SERVISKODU", output_dir / "02-serviskodu-hub-highres"
        )
    )
    generated.extend(draw_hub_star(buckets["DPK"], "DPK", output_dir / "03-dpk-hub-highres"))
    generated.extend(
        draw_other_relationships(buckets["OTHER"], output_dir / "04-diger-iliskiler-highres")
    )

    full_csv = output_dir / "99-iliskiler-tam.csv"
    write_relationship_csv(relationships, full_csv)
    write_relationship_csv(buckets["STATU"], output_dir / "statu-iliskileri.csv")
    write_relationship_csv(buckets["SERVISKODU"], output_dir / "serviskodu-iliskileri.csv")
    write_relationship_csv(buckets["DPK"], output_dir / "dpk-iliskileri.csv")
    write_relationship_csv(buckets["OTHER"], output_dir / "diger-iliskiler.csv")

    readme = write_readme(output_dir, tables, relationships, buckets)

    print(f"Input: {input_xlsx}")
    print(f"Output dir: {output_dir}")
    print(f"Tables: {len(tables)}")
    print(f"Relationships: {len(relationships)}")
    for key in ("STATU", "SERVISKODU", "DPK", "OTHER"):
        print(f"{key}: {len(buckets[key])}")
    print("Generated files:")
    for p in generated:
        print(f"- {p}")
    print(f"- {full_csv}")
    print(f"- {readme}")


if __name__ == "__main__":
    main()
