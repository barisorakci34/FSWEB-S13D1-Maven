#!/usr/bin/env python3
"""Generate readable full-CRM ER diagram package from workbook."""

from __future__ import annotations

import argparse
import csv
import math
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Sequence, Tuple

import matplotlib.pyplot as plt
import networkx as nx
from openpyxl import load_workbook


@dataclass
class FkRow:
    constraint: str
    source_table: str
    source_columns: List[str]
    target_table: str
    target_in_script: bool


def short_name(full_name: str) -> str:
    return full_name.split(".", 1)[1] if "." in full_name else full_name


def slug(text: str) -> str:
    s = text.lower().replace(".", "-")
    s = re.sub(r"[^a-z0-9\\-]+", "-", s)
    return re.sub(r"-+", "-", s).strip("-")


def split_cols(value: str) -> List[str]:
    if not value:
        return []
    return [x.strip() for x in str(value).split(",") if x.strip()]


def summarize_columns(columns: Iterable[str], max_show: int = 2) -> str:
    vals: List[str] = []
    seen = set()
    for col in columns:
        key = col.lower()
        if key in seen:
            continue
        seen.add(key)
        vals.append(col)
    if not vals:
        return ""
    if len(vals) <= max_show:
        return ", ".join(vals)
    return ", ".join(vals[:max_show]) + f" +{len(vals)-max_show}"


def read_workbook(path: Path) -> Tuple[Dict[str, int], List[FkRow]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws_tables = wb["Tablolar"]
    ws_fk = wb["Iliskiler_FK"]

    table_col_counts: Dict[str, int] = {}
    for row in ws_tables.iter_rows(min_row=2, values_only=True):
        schema, table, col_count = row[0], row[1], row[2]
        if not (schema and table):
            continue
        table_col_counts[f"{schema}.{table}"] = int(col_count or 0)

    fk_rows: List[FkRow] = []
    for row in ws_fk.iter_rows(min_row=2, values_only=True):
        constraint, ss, st, sc, ts, tt, _tc, _da, _ua, tis = row
        if not (ss and st and ts and tt):
            continue
        fk_rows.append(
            FkRow(
                constraint=str(constraint or ""),
                source_table=f"{ss}.{st}",
                source_columns=split_cols(str(sc or "")),
                target_table=f"{ts}.{tt}",
                target_in_script=str(tis or "").strip().lower() == "yes",
            )
        )

    return table_col_counts, fk_rows


def aggregate_pairs(fk_rows: Sequence[FkRow]) -> Dict[Tuple[str, str], List[FkRow]]:
    pairs: Dict[Tuple[str, str], List[FkRow]] = defaultdict(list)
    for fk in fk_rows:
        pairs[(fk.source_table, fk.target_table)].append(fk)
    return dict(sorted(pairs.items(), key=lambda kv: (kv[0][0], kv[0][1])))


def build_graph(
    table_col_counts: Dict[str, int],
    pair_map: Dict[Tuple[str, str], List[FkRow]],
) -> nx.DiGraph:
    g = nx.DiGraph()
    all_nodes = set(table_col_counts.keys())
    for source, target in pair_map:
        all_nodes.add(source)
        all_nodes.add(target)

    for node in all_nodes:
        g.add_node(
            node,
            col_count=table_col_counts.get(node, 0),
            internal=node in table_col_counts,
        )

    for (source, target), rows in pair_map.items():
        src_cols = [c for r in rows for c in r.source_columns]
        g.add_edge(
            source,
            target,
            fk_count=len(rows),
            columns=summarize_columns(src_cols, max_show=2),
        )

    return g


def compute_layout(g: nx.DiGraph, seed: int = 42) -> Dict[str, Tuple[float, float]]:
    n = max(g.number_of_nodes(), 1)
    k = 1.55 / math.sqrt(n)
    return nx.spring_layout(g.to_undirected(), seed=seed, k=k, iterations=700)


def draw_overview(
    g: nx.DiGraph,
    output_base: Path,
    title: str,
    exclude_node: str | None = None,
) -> None:
    graph = g.copy()
    if exclude_node and exclude_node in graph:
        graph.remove_node(exclude_node)

    pos = compute_layout(graph, seed=43 if exclude_node else 41)
    fig, ax = plt.subplots(figsize=(30, 20))
    ax.set_axis_off()

    indeg = dict(graph.in_degree())
    outdeg = dict(graph.out_degree())
    degree = {n: indeg.get(n, 0) + outdeg.get(n, 0) for n in graph.nodes}

    node_sizes = []
    node_colors = []
    for n in graph.nodes:
        d = degree.get(n, 0)
        col_count = graph.nodes[n].get("col_count", 0)
        node_sizes.append(120 + min(2200, d * 36 + col_count * 0.45))
        if n == "Crm.Users":
            node_colors.append("#D9534F")
        elif d >= 40:
            node_colors.append("#F0AD4E")
        elif d >= 12:
            node_colors.append("#5BC0DE")
        else:
            node_colors.append("#9FBAD0")

    edge_widths = [0.35 + min(2.9, graph.edges[e]["fk_count"] * 0.22) for e in graph.edges]
    nx.draw_networkx_edges(
        graph,
        pos=pos,
        arrows=True,
        arrowstyle="-|>",
        arrowsize=9,
        width=edge_widths,
        alpha=0.20,
        edge_color="#2F3A45",
        connectionstyle="arc3,rad=0.03",
        ax=ax,
    )
    nx.draw_networkx_nodes(
        graph,
        pos=pos,
        node_size=node_sizes,
        node_color=node_colors,
        edgecolors="#1F2A35",
        linewidths=0.5,
        ax=ax,
    )

    max_labels = 65 if graph.number_of_nodes() > 150 else 45
    label_nodes = set(
        [n for n, _ in sorted(degree.items(), key=lambda kv: (-kv[1], kv[0]))[:max_labels]]
    )
    labels = {n: short_name(n) for n in label_nodes}
    nx.draw_networkx_labels(graph, pos=pos, labels=labels, font_size=6.2, font_color="#101418", ax=ax)

    ax.set_title(
        (
            f"{title}\n"
            f"Node: {graph.number_of_nodes()} | Aggregated FK edges: {graph.number_of_edges()} | "
            f"Labelled nodes: {len(labels)}"
        ),
        fontsize=14,
        pad=14,
    )

    legend_items = [
        plt.Line2D([0], [0], marker="o", color="w", label="High degree (>=40)", markerfacecolor="#F0AD4E", markeredgecolor="#1F2A35", markersize=8),
        plt.Line2D([0], [0], marker="o", color="w", label="Medium degree (>=12)", markerfacecolor="#5BC0DE", markeredgecolor="#1F2A35", markersize=8),
        plt.Line2D([0], [0], marker="o", color="w", label="Other nodes", markerfacecolor="#9FBAD0", markeredgecolor="#1F2A35", markersize=8),
        plt.Line2D([0], [0], color="#2F3A45", lw=2, label="Explicit FK"),
    ]
    ax.legend(handles=legend_items, loc="upper right", framealpha=0.92, fontsize=9)

    output_base.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output_base.with_suffix(".png"), dpi=420, format="png", bbox_inches="tight", pad_inches=0.22)
    fig.savefig(output_base.with_suffix(".svg"), format="svg", bbox_inches="tight", pad_inches=0.22)
    plt.close(fig)


def draw_hub_pages(
    hub_table: str,
    pair_map: Dict[Tuple[str, str], List[FkRow]],
    table_col_counts: Dict[str, int],
    output_dir: Path,
    page_size: int = 55,
) -> int:
    incoming: Dict[str, List[FkRow]] = defaultdict(list)
    for (source, target), rows in pair_map.items():
        if target == hub_table:
            incoming[source].extend(rows)

    sources = sorted(incoming.keys(), key=lambda s: (-len(incoming[s]), s))
    if not sources:
        return 0

    page_count = math.ceil(len(sources) / page_size)
    hub_slug = slug(hub_table)

    for page_idx in range(page_count):
        page_sources = sources[page_idx * page_size : (page_idx + 1) * page_size]
        graph = nx.DiGraph()
        graph.add_node(hub_table)
        for src in page_sources:
            rows = incoming[src]
            cols = [c for r in rows for c in r.source_columns]
            graph.add_node(src)
            graph.add_edge(
                src,
                hub_table,
                fk_count=len(rows),
                columns=summarize_columns(cols, max_show=3),
            )

        radius = max(6.5, len(page_sources) * 0.23)
        pos = {hub_table: (0.0, 0.0)}
        for i, node in enumerate(page_sources):
            angle = (2.0 * math.pi * i) / max(1, len(page_sources))
            pos[node] = (radius * math.cos(angle), radius * math.sin(angle))

        fig, ax = plt.subplots(figsize=(24, 24))
        ax.set_axis_off()

        node_sizes = [4200]
        node_colors = ["#D9534F" if hub_table == "Crm.Users" else "#2E3A46"]
        for src in page_sources:
            fk_count = graph.edges[(src, hub_table)]["fk_count"]
            col_count = table_col_counts.get(src, 0)
            node_sizes.append(700 + min(2200, fk_count * 110 + col_count * 0.9))
            node_colors.append("#7FB3D5")

        nodes_order = [hub_table] + page_sources
        nx.draw_networkx_nodes(
            graph,
            pos=pos,
            nodelist=nodes_order,
            node_size=node_sizes,
            node_color=node_colors,
            edgecolors="#1F2A35",
            linewidths=1.0,
            ax=ax,
        )
        nx.draw_networkx_edges(
            graph,
            pos=pos,
            arrows=True,
            arrowstyle="-|>",
            arrowsize=16,
            width=1.25,
            alpha=0.66,
            edge_color="#2F3A45",
            connectionstyle="arc3,rad=0.08",
            ax=ax,
        )

        labels = {n: short_name(n) for n in nodes_order}
        nx.draw_networkx_labels(graph, pos=pos, labels=labels, font_size=8.0, font_color="#101418", ax=ax)
        nx.draw_networkx_labels(
            graph,
            pos={hub_table: pos[hub_table]},
            labels={hub_table: short_name(hub_table)},
            font_size=11,
            font_color="#FFFFFF",
            font_weight="bold",
            ax=ax,
        )

        edge_labels = {}
        for src in page_sources:
            edge_labels[(src, hub_table)] = graph.edges[(src, hub_table)]["columns"]
        nx.draw_networkx_edge_labels(
            graph,
            pos=pos,
            edge_labels=edge_labels,
            font_size=6.4,
            rotate=False,
            bbox={"facecolor": "white", "edgecolor": "none", "alpha": 0.82, "pad": 0.15},
            ax=ax,
        )

        ax.set_title(
            (
                f"CRM Full ER Hub: {hub_table} (page {page_idx+1}/{page_count})\n"
                f"Page source nodes: {len(page_sources)} | Total source nodes: {len(sources)}"
            ),
            fontsize=14,
            pad=16,
        )

        out = output_dir / f"hub-{hub_slug}-p{page_idx+1}-highres"
        fig.savefig(out.with_suffix(".png"), dpi=420, format="png", bbox_inches="tight", pad_inches=0.2)
        fig.savefig(out.with_suffix(".svg"), format="svg", bbox_inches="tight", pad_inches=0.2)
        plt.close(fig)

    return page_count


def write_summary_files(
    output_dir: Path,
    pair_map: Dict[Tuple[str, str], List[FkRow]],
    top_hubs: Sequence[Tuple[str, int]],
    hub_pages: Dict[str, int],
    graph: nx.DiGraph,
) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)

    pair_csv = output_dir / "fk-pairs-summary.csv"
    with pair_csv.open("w", newline="", encoding="utf-8") as f:
        wr = csv.writer(f)
        wr.writerow(["source_table", "target_table", "fk_count", "sample_source_columns"])
        for (source, target), rows in pair_map.items():
            cols = [c for r in rows for c in r.source_columns]
            wr.writerow([source, target, len(rows), summarize_columns(cols, max_show=4)])

    hub_csv = output_dir / "hub-summary.csv"
    with hub_csv.open("w", newline="", encoding="utf-8") as f:
        wr = csv.writer(f)
        wr.writerow(["hub_table", "incoming_fk_count", "generated_pages"])
        for hub, count in top_hubs:
            wr.writerow([hub, count, hub_pages.get(hub, 0)])

    readme = output_dir / "README.md"
    lines = [
        "# CRM Full ER Diagram Package",
        "",
        "- Kaynak: `docs/crm-full-veri-sozlugu.xlsx`",
        f"- Toplam node: **{graph.number_of_nodes()}**",
        f"- Toplam aggregated edge: **{graph.number_of_edges()}**",
        "",
        "## Dosyalar",
        "",
        "- `00-overview-fk-highres.png/.svg`: Tum FK iliskileri genel gorunum",
        "- `01-overview-fk-without-users-highres.png/.svg`: `Crm.Users` cikarilmis gorunum",
        "- `hub-*.png/.svg`: En cok referans alan tablolar icin sayfali hub diyagramlari",
        "- `fk-pairs-summary.csv`: Source-target bazli FK ozeti",
        "- `hub-summary.csv`: Hub tablo sayfalama ozeti",
        "",
        "Not: Bu paket explicit FK'lerden uretildi; tahmini iliski yoktur.",
        "",
    ]
    readme.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate full CRM ER diagram package")
    parser.add_argument(
        "--input-xlsx",
        default="docs/crm-full-veri-sozlugu.xlsx",
        help="Input workbook path",
    )
    parser.add_argument(
        "--output-dir",
        default="docs/crm-full-er",
        help="Output directory",
    )
    parser.add_argument(
        "--top-hub-count",
        type=int,
        default=6,
        help="How many inbound hubs to generate",
    )
    parser.add_argument(
        "--hub-page-size",
        type=int,
        default=55,
        help="How many source nodes per hub page",
    )
    args = parser.parse_args()

    input_xlsx = Path(args.input_xlsx)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    table_col_counts, fk_rows = read_workbook(input_xlsx)
    pair_map = aggregate_pairs(fk_rows)
    graph = build_graph(table_col_counts, pair_map)

    draw_overview(
        g=graph,
        output_base=output_dir / "00-overview-fk-highres",
        title="CRM Full ER - Explicit FK Overview",
        exclude_node=None,
    )
    draw_overview(
        g=graph,
        output_base=output_dir / "01-overview-fk-without-users-highres",
        title="CRM Full ER - Overview (Crm.Users excluded)",
        exclude_node="Crm.Users",
    )

    incoming_counts = Counter(fk.target_table for fk in fk_rows)
    top_hubs = incoming_counts.most_common(args.top_hub_count)
    hub_pages: Dict[str, int] = {}
    for hub, _count in top_hubs:
        pages = draw_hub_pages(
            hub_table=hub,
            pair_map=pair_map,
            table_col_counts=table_col_counts,
            output_dir=output_dir,
            page_size=max(15, args.hub_page_size),
        )
        hub_pages[hub] = pages

    write_summary_files(
        output_dir=output_dir,
        pair_map=pair_map,
        top_hubs=top_hubs,
        hub_pages=hub_pages,
        graph=graph,
    )

    print(f"Input: {input_xlsx}")
    print(f"Output dir: {output_dir}")
    print(f"Nodes: {graph.number_of_nodes()}")
    print(f"Aggregated edges: {graph.number_of_edges()}")
    print(f"FK rows: {len(fk_rows)}")
    for hub, count in top_hubs:
        print(f"Hub {hub}: {count} FK, pages={hub_pages.get(hub, 0)}")


if __name__ == "__main__":
    main()
