#!/usr/bin/env python3
"""Generate a concise CRM critical-table package (xlsx + ER)."""

from __future__ import annotations

import argparse
import csv
import re
from collections import defaultdict
from pathlib import Path
from typing import Dict, Iterable, List, Sequence, Tuple

import matplotlib.pyplot as plt
import networkx as nx
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


CRITICAL_TABLES: List[str] = [
    "Crm.Users",
    "Crm.Organizations",
    "Crm.Contracts",
    "Crm.Customers",
    "Crm.Activities",
    "Crm.Areas",
    "Crm.Reps",
    "Crm.Opportunities",
    "Crm.Mortgages",
    "Crm.Files",
    "Crm.Premiums",
    "Crm.Roles",
    "Crm.Products",
    "Crm.Portfolios",
    "Crm.Leads",
    "Crm.MortgageTypes",
    "Crm.PremiumContracts",
    "Crm.PremiumItems",
    "Crm.ActivityTypes",
    "Crm.Simulations",
    "Crm.Tickets",
    "Crm.Campaigns",
    "Crm.CustomerSources",
    "Crm.Groups",
    "Crm.Inventories",
]

AUDIT_USER_COLUMN_RE = re.compile(r"^(CreatedBy|UpdatedBy|DeletedBy)(UserUid)?$", re.IGNORECASE)


def short_name(full_name: str) -> str:
    return full_name.split(".", 1)[1] if "." in full_name else full_name


def style_sheet(ws, widths: Sequence[int]) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    center = Alignment(horizontal="center", vertical="center")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = center
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def read_full_workbook(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws_tables = wb["Tablolar"]
    ws_columns = wb["Kolonlar"]
    ws_fk = wb["Iliskiler_FK"]

    tables_rows: Dict[str, Tuple] = {}
    for row in ws_tables.iter_rows(min_row=2, values_only=True):
        schema, table = row[0], row[1]
        if schema and table:
            tables_rows[f"{schema}.{table}"] = row

    columns_rows: Dict[str, List[Tuple]] = defaultdict(list)
    for row in ws_columns.iter_rows(min_row=2, values_only=True):
        schema, table = row[0], row[1]
        if schema and table:
            columns_rows[f"{schema}.{table}"].append(row)

    fk_rows: List[Tuple] = []
    for row in ws_fk.iter_rows(min_row=2, values_only=True):
        if row[1] and row[2] and row[4] and row[5]:
            fk_rows.append(row)

    return tables_rows, columns_rows, fk_rows


def filter_fk(
    fk_rows: Sequence[Tuple],
    selected: Sequence[str],
) -> Tuple[List[Tuple], List[Tuple]]:
    selected_set = set(selected)
    fk_inside = []
    fk_external = []
    for row in fk_rows:
        source = f"{row[1]}.{row[2]}"
        target = f"{row[4]}.{row[5]}"
        if source not in selected_set:
            continue
        if target in selected_set:
            fk_inside.append(row)
        else:
            fk_external.append(row)
    return fk_inside, fk_external


def write_excel(
    output_xlsx: Path,
    selected: Sequence[str],
    tables_rows: Dict[str, Tuple],
    columns_rows: Dict[str, List[Tuple]],
    fk_inside: Sequence[Tuple],
    fk_external: Sequence[Tuple],
) -> None:
    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "Ozet"
    ws_tables = wb.create_sheet("Tablolar")
    ws_columns = wb.create_sheet("Kolonlar")
    ws_fk_in = wb.create_sheet("FK_Secili_Ici")
    ws_fk_out = wb.create_sheet("FK_Secili_Disi")

    total_columns = sum(len(columns_rows.get(t, [])) for t in selected)
    ws_summary.append(["Alan", "Deger"])
    ws_summary.append(["Kritik Tablo Sayisi", len(selected)])
    ws_summary.append(["Toplam Kolon", total_columns])
    ws_summary.append(["FK (Secili tablolar arasi)", len(fk_inside)])
    ws_summary.append(["FK (Secili tablolardan disariya)", len(fk_external)])
    ws_summary.append(["Secili tablolar", ", ".join(selected)])
    style_sheet(ws_summary, [34, 160])

    ws_tables.append(
        [
            "Schema",
            "Table",
            "ColumnCount",
            "PKConstraint",
            "PKColumns",
            "UniqueConstraintCount",
            "DefaultColumnCount",
            "FKOutCount",
            "FKInCount",
        ]
    )
    for t in selected:
        row = tables_rows.get(t)
        if not row:
            continue
        ws_tables.append(list(row))
    style_sheet(ws_tables, [12, 30, 12, 34, 28, 22, 18, 12, 12])

    ws_columns.append(
        [
            "Schema",
            "Table",
            "Ordinal",
            "Column",
            "DataType",
            "Nullability",
            "IsPrimaryKey",
            "IsUniqueMember",
            "IsComputed",
            "ComputedExpression",
            "IsMasked",
            "IsIdentity",
            "DefaultExpression",
            "DefaultConstraint",
        ]
    )
    for t in selected:
        for row in columns_rows.get(t, []):
            ws_columns.append(list(row))
    style_sheet(ws_columns, [12, 24, 10, 32, 22, 12, 12, 13, 11, 56, 10, 10, 34, 34])

    fk_headers = [
        "Constraint",
        "SourceSchema",
        "SourceTable",
        "SourceColumns",
        "TargetSchema",
        "TargetTable",
        "TargetColumns",
        "DeleteAction",
        "UpdateAction",
        "TargetInScript",
    ]
    ws_fk_in.append(fk_headers)
    for row in fk_inside:
        ws_fk_in.append(list(row))
    style_sheet(ws_fk_in, [48, 14, 26, 26, 14, 26, 26, 16, 16, 14])

    ws_fk_out.append(fk_headers)
    for row in fk_external:
        ws_fk_out.append(list(row))
    style_sheet(ws_fk_out, [48, 14, 26, 26, 14, 26, 26, 16, 16, 14])

    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_xlsx)


def aggregate_pairs(rows: Sequence[Tuple]) -> Dict[Tuple[str, str], Dict[str, object]]:
    pairs: Dict[Tuple[str, str], Dict[str, object]] = {}
    for row in rows:
        source = f"{row[1]}.{row[2]}"
        target = f"{row[4]}.{row[5]}"
        cols = [c.strip() for c in str(row[3] or "").split(",") if c.strip()]
        key = (source, target)
        if key not in pairs:
            pairs[key] = {"count": 0, "columns": []}
        pairs[key]["count"] = int(pairs[key]["count"]) + 1
        pairs[key]["columns"].extend(cols)
    return pairs


def is_audit_user_row(row: Tuple) -> bool:
    target = f"{row[4]}.{row[5]}"
    if target != "Crm.Users":
        return False
    cols = [c.strip() for c in str(row[3] or "").split(",") if c.strip()]
    if not cols:
        return False
    return all(AUDIT_USER_COLUMN_RE.match(c) for c in cols)


def summarize_columns(columns: Iterable[str], max_show: int = 2) -> str:
    uniq = []
    seen = set()
    for c in columns:
        k = c.lower()
        if k in seen:
            continue
        seen.add(k)
        uniq.append(c)
    if not uniq:
        return ""
    if len(uniq) <= max_show:
        return ", ".join(uniq)
    return ", ".join(uniq[:max_show]) + f" +{len(uniq)-max_show}"


def draw_er(
    selected: Sequence[str],
    fk_inside: Sequence[Tuple],
    output_base: Path,
    title: str,
) -> None:
    pair_map = aggregate_pairs(fk_inside)
    g = nx.DiGraph()
    for t in selected:
        g.add_node(t)
    for (s, t), payload in pair_map.items():
        g.add_edge(s, t, fk_count=int(payload["count"]), cols=summarize_columns(payload["columns"]))

    pos = nx.spring_layout(g.to_undirected(), seed=42, k=0.95, iterations=700)
    fig, ax = plt.subplots(figsize=(22, 15))
    ax.set_axis_off()

    indeg = dict(g.in_degree())
    outdeg = dict(g.out_degree())
    degree = {n: indeg.get(n, 0) + outdeg.get(n, 0) for n in g.nodes}

    sizes = [1200 + degree[n] * 140 for n in g.nodes]
    colors = []
    for n in g.nodes:
        if n == "Crm.Users":
            colors.append("#D9534F")
        elif degree[n] >= 8:
            colors.append("#F0AD4E")
        else:
            colors.append("#5BC0DE")

    widths = [1.1 + min(3.5, g.edges[e]["fk_count"] * 0.55) for e in g.edges]
    nx.draw_networkx_edges(
        g,
        pos=pos,
        arrows=True,
        arrowstyle="-|>",
        arrowsize=16,
        width=widths,
        alpha=0.7,
        edge_color="#2F3A45",
        connectionstyle="arc3,rad=0.08",
        ax=ax,
    )
    nx.draw_networkx_nodes(
        g,
        pos=pos,
        node_size=sizes,
        node_color=colors,
        edgecolors="#1F2A35",
        linewidths=1.0,
        ax=ax,
    )
    labels = {n: short_name(n) for n in g.nodes}
    nx.draw_networkx_labels(g, pos=pos, labels=labels, font_size=8.2, font_color="#101418", ax=ax)

    edge_labels = {(u, v): d["cols"] for u, v, d in g.edges(data=True) if d.get("cols")}
    if edge_labels:
        nx.draw_networkx_edge_labels(
            g,
            pos=pos,
            edge_labels=edge_labels,
            font_size=7.0,
            rotate=False,
            bbox={"facecolor": "white", "edgecolor": "none", "alpha": 0.78, "pad": 0.18},
            ax=ax,
        )

    ax.set_title((f"{title}\nTablo: {g.number_of_nodes()} | Iliski (aggregate): {g.number_of_edges()}"), fontsize=15, pad=14)

    output_base.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output_base.with_suffix(".png"), dpi=420, format="png", bbox_inches="tight", pad_inches=0.22)
    fig.savefig(output_base.with_suffix(".svg"), format="svg", bbox_inches="tight", pad_inches=0.22)
    plt.close(fig)


def write_csvs(
    output_dir: Path,
    fk_inside: Sequence[Tuple],
    fk_external: Sequence[Tuple],
) -> None:
    headers = [
        "Constraint",
        "SourceSchema",
        "SourceTable",
        "SourceColumns",
        "TargetSchema",
        "TargetTable",
        "TargetColumns",
        "DeleteAction",
        "UpdateAction",
        "TargetInScript",
    ]

    def _write(path: Path, rows: Sequence[Tuple]) -> None:
        with path.open("w", newline="", encoding="utf-8") as f:
            wr = csv.writer(f)
            wr.writerow(headers)
            wr.writerows(rows)

    output_dir.mkdir(parents=True, exist_ok=True)
    _write(output_dir / "fk-secili-ici.csv", fk_inside)
    _write(output_dir / "fk-secili-disi.csv", fk_external)


def write_readme(output_dir: Path, selected: Sequence[str], fk_inside: Sequence[Tuple], fk_external: Sequence[Tuple]) -> None:
    p = output_dir / "README.md"
    lines = [
        "# CRM Kritik 25 Tablo Paketi",
        "",
        "Bu paket sadece kritik 25 tabloyu ve dogrulanmis FK iliskilerini icerir.",
        "",
        f"- Kritik tablo sayisi: **{len(selected)}**",
        f"- Secili tablolar arasi FK: **{len(fk_inside)}**",
        f"- Secili tablolardan disari FK: **{len(fk_external)}**",
        "",
        "## Dosyalar",
        "",
        "- `crm-kritik-25-veri-sozlugu.xlsx`",
        "- `crm-kritik-25-er.png/.svg` (tum dogrulanmis FK)",
        "- `crm-kritik-25-er-audit-haric.png/.svg` (Users audit FK'lari haric)",
        "- `fk-secili-ici.csv`",
        "- `fk-secili-disi.csv`",
        "",
        "## Kritik Tablo Listesi",
        "",
    ]
    for t in selected:
        lines.append(f"- {t}")
    lines.append("")
    p.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate concise critical CRM package")
    parser.add_argument(
        "--input-xlsx",
        default="docs/crm-full-veri-sozlugu.xlsx",
        help="Input full dictionary workbook",
    )
    parser.add_argument(
        "--output-dir",
        default="docs/crm-kritik-25",
        help="Output directory",
    )
    args = parser.parse_args()

    input_xlsx = Path(args.input_xlsx)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    selected = CRITICAL_TABLES
    tables_rows, columns_rows, fk_rows = read_full_workbook(input_xlsx)
    fk_inside, fk_external = filter_fk(fk_rows, selected)

    write_excel(
        output_xlsx=output_dir / "crm-kritik-25-veri-sozlugu.xlsx",
        selected=selected,
        tables_rows=tables_rows,
        columns_rows=columns_rows,
        fk_inside=fk_inside,
        fk_external=fk_external,
    )
    draw_er(
        selected=selected,
        fk_inside=fk_inside,
        output_base=output_dir / "crm-kritik-25-er",
        title="CRM Kritik 25 ER (Tum Dogrulanmis FK)",
    )

    fk_inside_no_audit = [r for r in fk_inside if not is_audit_user_row(r)]
    draw_er(
        selected=selected,
        fk_inside=fk_inside_no_audit,
        output_base=output_dir / "crm-kritik-25-er-audit-haric",
        title="CRM Kritik 25 ER (Audit Users FK Haric)",
    )
    write_csvs(output_dir=output_dir, fk_inside=fk_inside, fk_external=fk_external)
    write_readme(output_dir=output_dir, selected=selected, fk_inside=fk_inside, fk_external=fk_external)

    print(f"Input: {input_xlsx}")
    print(f"Output dir: {output_dir}")
    print(f"Selected tables: {len(selected)}")
    print(f"FK inside: {len(fk_inside)}")
    print(f"FK external: {len(fk_external)}")


if __name__ == "__main__":
    main()
